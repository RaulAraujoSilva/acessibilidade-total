"""
Gera tests/deck_bom.pptx — o deck de controle.

Serve ao teste inverso: um auditor que acusa tudo e tao inutil quanto um que
nao acusa nada. Aqui NAO pode sobrar Erro nem Aviso; o que aparecer e falso
positivo do auditor.

E feito A MAO de proposito, sem passar pelo build_deck. Se o controle fosse a
saida do proprio construtor, o teste so provaria que os dois concordam entre
si — e nao que ambos estao certos. O que ele compartilha com o construtor sao
as MEDIDAS (scripts/grade.py) e o modelo, que sao a definicao de "correto".

    python tests/make_deck_bom.py
    python tests/test_falso_positivo.py
"""
from __future__ import annotations

import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(RAIZ, "scripts"))

import a11y_lib as A  # noqa: E402
import grade as G  # noqa: E402

SAIDA = os.path.join(HERE, "deck_bom.pptx")
MODELO = os.path.join(RAIZ, "assets", "modelo-acessivel.pptx")

FUNDO = "FAF7F2"
TEXTO = RGBColor(0x1A, 0x1A, 0x1A)
DESTAQUE = RGBColor(0x00, 0x72, 0xB2)


def png_solido(path, cor=(0, 114, 178), w=900, h=600):
    import struct
    import zlib
    raw = b"".join(b"\x00" + bytes(cor) * w for _ in range(h))

    def chunk(t, d):
        c = t + d
        return struct.pack(">I", len(d)) + c + struct.pack(">I", zlib.crc32(c))

    with open(path, "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\n"
                + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(raw, 9))
                + chunk(b"IEND", b""))
    return path, w, h


def pinta_fundo(slide):
    cSld = slide._element.find(A.q("p:cSld"))
    antigo = cSld.find(A.q("p:bg"))
    if antigo is not None:
        cSld.remove(antigo)
    bg = etree.Element(A.q("p:bg"))
    bgPr = etree.SubElement(bg, A.q("p:bgPr"))
    fill = etree.SubElement(bgPr, A.q("a:solidFill"))
    etree.SubElement(fill, A.q("a:srgbClr")).set("val", FUNDO)
    etree.SubElement(bgPr, A.q("a:effectLst"))
    cSld.insert(0, bg)


def formata(tf, tamanho=24, cor=None, negrito=False, entrelinha=None):
    ln_val = entrelinha or G.ENTRELINHA_CORPO
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        pPr = p._p.get_or_add_pPr()
        for antigo in pPr.findall(A.q("a:lnSpc")):
            pPr.remove(antigo)
        ln = etree.Element(A.q("a:lnSpc"))
        etree.SubElement(ln, A.q("a:spcPct")).set("val", str(ln_val))
        pPr.insert(0, ln)
        for r in p.runs:
            r.font.size = Pt(tamanho)
            r.font.name = "Calibri"
            r.font.bold = negrito
            r.font.color.rgb = cor or TEXTO
            rPr = r._r.find(A.q("a:rPr"))
            if rPr is not None:
                rPr.set("lang", "pt-BR")


def nomeia(shape, nome):
    c = A.get_cNvPr(shape._element)
    if c is not None:
        c.set("name", nome)


def poe(shape, coluna, n, y, h):
    shape.left, shape.top = G.x(coluna), y
    shape.width, shape.height = G.larg(n), h


def layout(prs, nome):
    for lay in prs.slide_layouts:
        if (lay.name or "").strip().lower() == nome.lower():
            return lay
    raise RuntimeError("modelo sem o layout %r — rode gerar_modelo.py" % nome)


def faixa_rodape(slide):
    f = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, G.ALTURA - G.FAIXA_H,
                               G.LARGURA, G.FAIXA_H)
    f.fill.solid()
    f.fill.fore_color.rgb = DESTAQUE
    f.line.fill.background()
    f.shadow.inherit = False
    nomeia(f, "Faixa estética do rodapé")
    A.set_decorative(f._element, True)
    # decorativo vai para o fundo do eixo Z
    spTree = slide.shapes._spTree
    spTree.remove(f._element)
    spTree.insert(2, f._element)


def main():
    if not os.path.exists(MODELO):
        print("modelo ausente — rode antes: python scripts/gerar_modelo.py")
        return 1

    prs = Presentation(MODELO)
    prs.slide_width, prs.slide_height = G.LARGURA, G.ALTURA
    img, nw, nh = png_solido(os.path.join(HERE, "_fig_bom.png"))

    # -- slide 1: texto e figura -------------------------------------------
    s1 = prs.slides.add_slide(layout(prs, "Título e conteúdo"))
    pinta_fundo(s1)
    t = s1.shapes.title
    t.text = "Por que a estrutura vem antes da estética"
    nomeia(t, "Título do slide 1")
    poe(t, 0, 12, G.TITULO_Y, G.TITULO_H)
    formata(t.text_frame, 34, negrito=True, entrelinha=G.ENTRELINHA_TITULO)

    corpo = s1.placeholders[1]
    nomeia(corpo, "Argumento principal")
    poe(corpo, 0, 5, G.CONTEUDO_Y, G.CONTEUDO_H)
    tf = corpo.text_frame
    tf.word_wrap = True
    tf.text = "O leitor de tela navega pela hierarquia, não pelo desenho"
    tf.add_paragraph().text = "Placeholder carrega semântica; caixa solta não"
    tf.add_paragraph().text = "Composição também se audita: proporção, margem, ordem"
    formata(tf, 24)

    # figura ENCAIXADA, sem esticar (regra N01)
    caixa = G.bloco(6, 6, G.CONTEUDO_Y, G.CONTEUDO_H)
    w, h, dx, dy = G.encaixar(nw, nh, caixa[2], caixa[3])
    pic = s1.shapes.add_picture(img, caixa[0] + dx, caixa[1] + dy, w, h)
    nomeia(pic, "Diagrama da hierarquia")
    A.set_alt_text(pic._element, "Hierarquia de três níveis, do título ao "
                                 "conteúdo e às notas do orador")
    s1.notes_slide.notes_text_frame.text = (
        "Descrição da figura: três blocos empilhados representando título, "
        "conteúdo e notas do orador, ligados de cima para baixo.")
    faixa_rodape(s1)

    # -- slide 2: tabela e link --------------------------------------------
    s2 = prs.slides.add_slide(layout(prs, "Título e conteúdo"))
    pinta_fundo(s2)
    t2 = s2.shapes.title
    t2.text = "O que o verificador nativo não enxerga"
    nomeia(t2, "Título do slide 2")
    poe(t2, 0, 12, G.TITULO_Y, G.TITULO_H)
    formata(t2.text_frame, 34, negrito=True, entrelinha=G.ENTRELINHA_TITULO)

    ref = s2.placeholders[1]
    nomeia(ref, "Referência normativa")
    poe(ref, 0, 12, G.CONTEUDO_Y, G.cm(1.5))
    r = ref.text_frame.paragraphs[0].add_run()
    r.text = "Diretrizes WCAG 2.2 do W3C"
    r.hyperlink.address = "https://www.w3.org/TR/WCAG22/"
    formata(ref.text_frame, 24, cor=DESTAQUE)
    r.font.underline = True

    linhas = [["Camada", "Cobertura nativa"],
              ["Texto alternativo", "Presença, não qualidade"],
              ["Tipografia", "Nenhuma"],
              ["Composição", "Nenhuma"],
              ["PDF exportado", "Nenhuma"]]
    y = G.CONTEUDO_Y + G.cm(2.2)
    linha_h = (G.RODAPE_Y - y) // len(linhas)
    gf = s2.shapes.add_table(len(linhas), 2, G.x(0), y, G.larg(12),
                             linha_h * len(linhas))
    nomeia(gf, "Tabela de cobertura")
    A.set_alt_text(gf._element, "Comparação entre o que o verificador nativo "
                                "cobre e o que fica de fora")
    tb = gf.table
    tb.first_row = True
    tblPr = tb._tbl.find(A.q("a:tblPr"))
    if tblPr is not None:
        tblPr.set("bandRow", "0")

    def pinta(cel, hexv):
        tcPr = cel._tc.get_or_add_tcPr()
        fill = etree.SubElement(tcPr, A.q("a:solidFill"))
        etree.SubElement(fill, A.q("a:srgbClr")).set("val", hexv)
        tcPr.remove(fill)
        tcPr.insert(0, fill)

    for ri, linha in enumerate(linhas):
        for ci, v in enumerate(linha):
            cel = tb.cell(ri, ci)
            cel.text = v
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            pinta(cel, "0072B2" if ri == 0 else "FFFFFF")
            formata(cel.text_frame, 20, entrelinha=100000)
            if ri == 0:
                for p_ in cel.text_frame.paragraphs:
                    for r_ in p_.runs:
                        r_.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        r_.font.bold = True
    faixa_rodape(s2)

    cp = prs.core_properties
    cp.title = "Acessibilidade estrutural em apresentações"
    cp.author = "Raul Araujo Silva"
    cp.language = "pt-BR"

    prs.save(SAIDA)
    os.remove(img)
    print("gerado:", SAIDA)


if __name__ == "__main__":
    main()
