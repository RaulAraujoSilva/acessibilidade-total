"""
Gera tests/deck_ruim.pptx — o deck-armadilha.

Cada defeito plantado aqui tem a regra do catalogo que DEVE acusa-lo anotada
em GABARITO. Se o auditor nao pegar um deles, o errado e o auditor, nao o deck.

    python tests/make_deck_ruim.py
    python tests/test_auditor.py
"""
from __future__ import annotations

import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))

import a11y_lib as A  # noqa: E402

SAIDA = os.path.join(HERE, "deck_ruim.pptx")

# Regras que o auditor tem obrigacao de acusar neste arquivo.
GABARITO = [
    "A01", "A02", "A03", "A04",
    "B01", "B03", "B05", "B06", "B09",
    "C01", "C02", "C05",
    "D01", "D03", "D04", "D05", "D06", "D08", "D11",
    "E01", "E02", "E08",
    "F01", "F02", "F04", "F06",
    "G01", "G03", "G04", "G05",
    "H01", "H02",
    "I07",
    # camada N — defeitos de composicao
    "N01", "N02", "N04", "N09",
]


def png_solido(path, cor=(200, 60, 60), w=240, h=160):
    """PNG minimo, sem depender de Pillow."""
    import struct
    import zlib

    raw = b"".join(b"\x00" + bytes(cor) * w for _ in range(h))

    def chunk(tipo, dados):
        c = tipo + dados
        return struct.pack(">I", len(dados)) + c + struct.pack(">I", zlib.crc32(c))

    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(raw, 9))
           + chunk(b"IEND", b""))
    with open(path, "wb") as f:
        f.write(png)
    return path


def set_lang(shape, lang="en-US"):
    for p_el in A.iter_paragraphs(shape._element):
        for r_el in A.iter_runs(p_el):
            rPr = r_el.find(A.q("a:rPr"))
            if rPr is None:
                rPr = etree.SubElement(r_el, A.q("a:rPr"))
                r_el.remove(rPr)
                r_el.insert(0, rPr)
            rPr.set("lang", lang)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    img = png_solido(os.path.join(HERE, "_fig.png"))

    branco = RGBColor(0xFF, 0xFF, 0xFF)

    # -- slide 1: titulo duplicado, alt ruins, contraste, tipografia --------
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Resultados"                      # B03 (repete no 2)
    corpo = s1.placeholders[1]
    tf = corpo.text_frame
    tf.text = "- item digitado com hifen"                    # B09
    p = tf.add_paragraph()
    p.text = "TEXTO INTEIRO EM CAIXA ALTA NESTE SLIDE"       # F06
    p2 = tf.add_paragraph()
    p2.text = "Texto pequeno demais e justificado nesta linha de exemplo"
    p2.alignment = 4                                          # PP_ALIGN.JUSTIFY -> F04
    for run in p2.runs:
        run.font.size = Pt(12)                                # F02
        run.font.name = "Times New Roman"                     # F01
        run.font.italic = True                                # F06
    set_lang(corpo, "en-US")                                  # A03
    set_lang(s1.shapes.title, "en-US")

    # imagem sem alt text (D01)
    s1.shapes.add_picture(img, Inches(9), Inches(1.5), Inches(3), Inches(2))

    # -- slide 2: titulo repetido, ordem de leitura invertida ---------------
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    # conteudo inserido ANTES do titulo receber texto: o titulo continua sendo
    # o 1o no spTree, entao movemos ele para o fim de proposito (C01/C02)
    cx = s2.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(6), Inches(1.5))
    cx.text_frame.text = "Bloco de conteudo em caixa de texto solta"   # B05
    s2.shapes.title.text = "Resultados"                                 # B03
    spTree = s2.shapes._spTree
    titulo_el = s2.shapes.title._element
    spTree.remove(titulo_el)
    spTree.append(titulo_el)                                            # C01 + C02

    # imagem com alt = nome de arquivo (D03) e prefixo (D04)
    pic = s2.shapes.add_picture(img, Inches(8.5), Inches(1.2), Inches(3.5), Inches(2.2))
    A.set_alt_text(pic._element, "grafico1.png")                        # D03

    pic2 = s2.shapes.add_picture(img, Inches(8.5), Inches(3.8), Inches(3.5), Inches(2.2))
    A.set_alt_text(pic2._element, "Imagem de um grafico de barras com linhas azuis. "
                                  "Descricao gerada automaticamente")   # D04 + D05

    # -- slide 3: sem titulo, decorativo contraditorio, alt longo -----------
    s3 = prs.slides.add_slide(prs.slide_layouts[6])                     # layout Em Branco
    cx3 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    cx3.text_frame.text = "Slide sem titulo nenhum"                     # B01 (+B05)

    pic3 = s3.shapes.add_picture(img, Inches(1), Inches(2.5), Inches(3), Inches(2))
    A.set_alt_text(pic3._element, "Descricao muito longa que passa dos cento e cinquenta "
                                  "caracteres de proposito para acionar a regra de "
                                  "truncamento em leitores de tela mais antigos e "
                                  "assim validar o auditor")            # D06

    # decorativo E com descr ao mesmo tempo (D11) — escrito a mao
    pic4 = s3.shapes.add_picture(img, Inches(5), Inches(2.5), Inches(3), Inches(2))
    A.set_decorative(pic4._element, True)
    A.get_cNvPr(pic4._element).set("descr", "isto nao deveria coexistir")

    # contraste ruim: cinza claro sobre branco (E01) e preto puro sobre branco (E08)
    # 20pt = texto GRANDE -> a regra correta e E02 (exige 3:1), nao E01
    cx4 = s3.shapes.add_textbox(Inches(1), Inches(5), Inches(5), Inches(0.8))
    cx4.text_frame.text = "Texto grande cinza claro de baixo contraste"
    for run in cx4.text_frame.paragraphs[0].runs:
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)                 # E02

    # 16pt = texto NORMAL -> E01 (exige 4,5:1); tambem cai em F02 (<18pt)
    cx4b = s3.shapes.add_textbox(Inches(1), Inches(5.9), Inches(5), Inches(0.8))
    cx4b.text_frame.text = "Texto normal cinza de baixo contraste"
    for run in cx4b.text_frame.paragraphs[0].runs:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)                 # E01 + F02

    cx5 = s3.shapes.add_textbox(Inches(7), Inches(5), Inches(5), Inches(0.8))
    cx5.text_frame.text = "Preto puro sobre branco puro"
    for run in cx5.text_frame.paragraphs[0].runs:
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)                 # E08

    # -- slide 4: tabela problematica e links ruins -------------------------
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Tabela e links"

    gf = s4.shapes.add_table(4, 4, Inches(0.8), Inches(1.8), Inches(8), Inches(2.6))
    tbl = gf.table
    tbl.first_row = False                                               # G01
    dados = [["Mes", "Meta", "", "Real"],
             ["Jan", "100", "", "92"],
             ["Fev", "110", "", "97"],
             ["Mar", "120", "", "88"]]
    for ri, linha in enumerate(dados):
        for ci, v in enumerate(linha):
            tbl.cell(ri, ci).text = v                                   # coluna 2 vazia -> G04
    tbl.cell(0, 0).merge(tbl.cell(0, 1))                                # G03
    # sem alt text na tabela -> G05

    lk1 = s4.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(6), Inches(0.5))
    r1 = lk1.text_frame.paragraphs[0].add_run()
    r1.text = "https://www.w3.org/TR/WCAG22/"                           # H01
    r1.hyperlink.address = "https://www.w3.org/TR/WCAG22/"
    r1.font.size = Pt(20)

    lk2 = s4.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(6), Inches(0.5))
    r2 = lk2.text_frame.paragraphs[0].add_run()
    r2.text = "clique aqui"                                             # H02
    r2.hyperlink.address = "https://exemplo.org/pagina"
    r2.font.size = Pt(20)

    # alvo de clique minusculo com hiperlink (I07)
    mini = s4.shapes.add_textbox(Emu(200000), Emu(200000), Emu(150000), Emu(150000))
    rm = mini.text_frame.paragraphs[0].add_run()
    rm.text = "x"
    rm.hyperlink.address = "https://exemplo.org/mini"

    # -- slide 5: defeitos de COMPOSICAO (camada N) -------------------------
    s5 = prs.slides.add_slide(prs.slide_layouts[2])   # Section Header do Office
    s5.shapes.title.text = "Seção com defeitos de composição"   # N09: cap=all
    formata_titulo = s5.shapes.title.text_frame
    for p_ in formata_titulo.paragraphs:
        for r_ in p_.runs:
            r_.font.size = Pt(36)
    s5.placeholders[1].text_frame.text = "Este apoio fica ACIMA do título"  # N04

    # figura esticada de proposito: 3:2 nativa forcada em 1:2 (N01)
    img2 = png_solido(os.path.join(HERE, "_fig_larga.png"), (0, 114, 178), 300, 200)
    s5.shapes.add_picture(img2, Inches(8.5), Inches(0.6), Inches(2.0), Inches(4.0))

    # bloco invadindo a faixa do rodape (N02)
    baixo = s5.shapes.add_textbox(Inches(0.9), Inches(6.9), Inches(8), Inches(0.9))
    baixo.text_frame.text = "Este bloco entra na faixa reservada do rodapé"
    for p_ in baixo.text_frame.paragraphs:
        for r_ in p_.runs:
            r_.font.size = Pt(20)

    # -- metadados propositalmente ruins ------------------------------------
    cp = prs.core_properties
    cp.title = ""                                                       # A01
    cp.author = "Usuario do Windows"                                    # A04
    cp.language = ""                                                    # A02

    prs.save(SAIDA)
    os.remove(img)
    for extra in (os.path.join(HERE, '_fig_larga.png'),):
        if os.path.exists(extra):
            os.remove(extra)
    print("gerado:", SAIDA)
    print("gabarito com %d regras:" % len(GABARITO), ", ".join(GABARITO))


if __name__ == "__main__":
    main()
