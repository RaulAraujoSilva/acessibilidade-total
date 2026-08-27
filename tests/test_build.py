"""
Teste de integracao: o construtor tem de produzir algo que o AUDITOR aprove.

E o fecho do ciclo. De nada adianta um construtor que gera slide bonito e um
auditor que reprova o resultado: um dos dois esta errado.

    python tests/test_build.py
"""
from __future__ import annotations

import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "scripts"))

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from audit_pptx import audit, NAO_CONFORME  # noqa: E402
import audit_pacote  # noqa: E402
from build_deck import (ErroDeRoteiro, MODOS, construir,  # noqa: E402
                        recursos_do_perfil,
                        construir_conjunto)

EXEMPLO = os.path.join(ROOT, "exemplos", "roteiro-exemplo.yaml")

# Usado quando PyYAML nao esta instalado (a CI instala so python-pptx).
ROTEIRO_MINIMO = {
    "apresentacao": {
        "titulo": "Deck de integração do construtor",
        "autor": "Suíte de testes",
        "idioma": "pt-BR",
    },
    "slides": [
        {"tipo": "capa", "titulo": "Capa do teste de integração",
         "subtitulo": "Gerado por build_deck"},
        {"tipo": "secao", "titulo": "Seção de exemplo",
         "conteudo": ["Cabeçalho de seção com uma linha de apoio."]},
        {"titulo": "Slide de texto simples",
         "conteudo": ["Primeira linha do corpo.", "Segunda linha do corpo."],
         "notas": "Notas do orador."},
        {"titulo": "Slide com tabela de dados",
         "conteudo": ["Tabela só para dados."],
         "tabela": {"alt": "Limiares de contraste por tamanho de texto",
                    "cabecalho": ["Texto", "AA"],
                    "linhas": [["Normal", "4,5:1"], ["Grande", "3:1"]]}},
        {"titulo": "Slide com link descritivo",
         "conteudo": ["O texto do link diz o destino."],
         "links": [{"texto": "Diretrizes WCAG 2.2 do W3C",
                    "url": "https://www.w3.org/TR/WCAG22/"}]},
    ],
}


def carregar():
    try:
        import yaml  # noqa: F401
        from build_deck import carregar_roteiro
        return carregar_roteiro(EXEMPLO), "exemplos/roteiro-exemplo.yaml"
    except ImportError:
        return ROTEIRO_MINIMO, "roteiro embutido (PyYAML ausente)"


def auditar(caminho, rotulo):
    rep = audit(caminho)
    ruins = [f for f in rep.findings if f["veredito"] == NAO_CONFORME]
    c = rep.counts()
    print("  %-22s %d Erros, %d Avisos, %d Dicas"
          % (rotulo, c.get("E", 0), c.get("A", 0), c.get("D", 0)))
    for f in ruins:
        print("      %s %s | %s | %s"
              % (f["severidade"], f["regra"], f["onde"], f["detalhe"][:80]))
    return ruins


def testa_recusa(roteiro_ruim, esperado):
    """O construtor tem de PARAR antes de gerar um slide inacessivel."""
    try:
        with tempfile.TemporaryDirectory() as d:
            construir(roteiro_ruim, os.path.join(d, "x.pptx"))
    except ErroDeRoteiro as e:
        ok = esperado.lower() in str(e).lower()
        print("  %-22s %s" % (esperado, "recusou corretamente" if ok
                              else "recusou, mas por outro motivo: %s" % e))
        return ok
    print("  %-22s ACEITOU — deveria ter recusado" % esperado)
    return False


def testa_recursos_por_perfil(perfis, pasta) -> int:
    """
    O recurso segue o sentido que serve — e isso tem de ser VERIFICAVEL.

    Aqui nao ha COM nem PowerPoint: conta-se `ppt/media/` pelo zipfile. O que
    se testa e a regra O07 nos dois sentidos, porque um auditor que nunca acusa
    aprova qualquer coisa.
    """
    import zipfile

    falhas = 0
    for chave, caminho in sorted(perfis.items()):
        perfil = chave.split("/")[0]
        esperado = recursos_do_perfil(perfil)
        with zipfile.ZipFile(caminho) as z:
            midia = [n for n in z.namelist() if n.startswith("ppt/media/")]
        mp3 = sum(1 for n in midia if n.lower().endswith(".mp3"))
        rotulo = "recursos %s" % perfil
        # o deck recem-construido ainda nao tem midia embutida; o que se pode
        # afirmar aqui e que ele nao ganhou audio sozinho
        if mp3 and not esperado["audio"]:
            print("  %-28s ERRO: %d faixa(s) num perfil que dispensa áudio"
                  % (rotulo, mp3))
            falhas += 1
        else:
            print("  %-28s coerente com o mapa" % rotulo)

    rep = audit_pacote.auditar(list(perfis.values()), pasta)
    o07 = [f for f in rep.findings
           if f["regra"] == "O07" and f["veredito"] == audit_pacote.NAO_CONFORME]
    # o perfil libras ainda nao tem as janelas embutidas, entao O07 DEVE acusar
    pegou = any("janela de Libras por slide" in f["detalhe"] for f in o07)
    print("  %-28s %s" % ("O07 (sem as janelas)",
                          "acusou" if pegou else "ERRO: não acusou"))
    return falhas + (0 if pegou else 1)


def copy_roteiro_com_chave(roteiro):
    """Roteiro com mensagem_chave em todo slide que precisa de uma."""
    import copy

    r = copy.deepcopy(roteiro)
    for i, spec in enumerate(r["slides"], 1):
        if spec.get("tipo") in ("capa", "secao", "citacao"):
            continue
        spec["mensagem_chave"] = "Ideia central do slide %d." % i
    return r


def conta_slides(caminho, esperado, rotulo) -> int:
    from pptx import Presentation

    achado = len(Presentation(caminho).slides._sldIdLst)
    if achado == esperado:
        print("  %-28s %d slides" % (rotulo, achado))
        return 0
    print("  %-28s ERRO: %d slides, esperado %d"
          % (rotulo, achado, esperado))
    return 1


def testa_paridade(arquivos, pasta) -> int:
    """
    Camada O nos dois sentidos.

    O teste que importa e o SEGUNDO: um auditor de paridade que nunca acusa
    nada aprova qualquer coisa. Adultera-se uma copia e exige-se O01.
    """
    import shutil

    from pptx import Presentation

    rep = audit_pacote.auditar(arquivos, pasta)
    erros = [f for f in rep.findings
             if f["veredito"] == audit_pacote.NAO_CONFORME
             and f["regra"] in ("O01", "O02")]
    falhas = 0
    if erros:
        print("  %-28s ERRO: versoes divergentes sem motivo" % "paridade")
        falhas += 1
    else:
        print("  %-28s as 3 versoes batem" % "paridade")

    adulterado = os.path.join(pasta, "adulterado.pptx")
    shutil.copy(arquivos[1], adulterado)
    prs = Presentation(adulterado)
    for slide in prs.slides:
        mexeu = False
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.paragraphs[0].runs:
                sh.text_frame.paragraphs[0].runs[0].text = "DIVERGENCIA"
                mexeu = True
                break
        if mexeu:
            break
    prs.save(adulterado)

    rep2 = audit_pacote.auditar([arquivos[0], adulterado], pasta)
    pegou = any(f["regra"] == "O01" and f["veredito"] == audit_pacote.NAO_CONFORME
                for f in rep2.findings)
    print("  %-28s %s" % ("paridade (adulterada)",
                          "O01 acusou" if pegou else "ERRO: O01 nao acusou"))
    return falhas + (0 if pegou else 1)


def main() -> int:
    roteiro, fonte = carregar()
    print("=" * 72)
    print("Construtor -> auditor  (fonte: %s)" % fonte)
    print("=" * 72)

    falhas = 0
    esperado = len(roteiro["slides"])
    with tempfile.TemporaryDirectory() as d:
        simples = os.path.join(d, "deck_simples.pptx")
        construir(roteiro, simples)
        falhas += len(auditar(simples, "deck simples"))
        falhas += conta_slides(simples, esperado, "deck simples")

        # Um arquivo POR MODO, com o MESMO conteudo. A contagem de slides tem
        # de ser a do roteiro: quando os tres modos viviam num arquivo so, ela
        # era 1 + 3xN e ninguem verificava.
        feitos = construir_conjunto(roteiro, d, "conjunto")
        assert set(feitos) == {"completo/%s" % m for m in MODOS}, feitos
        for chave, caminho in feitos.items():
            falhas += len(auditar(caminho, "deck %s" % chave))
            falhas += conta_slides(caminho, esperado, "deck %s" % chave)


        # Perfis de publico. O roteiro de exemplo nao tem mensagem_chave, entao
        # a recusa E o comportamento certo: reduzir texto e trabalho de redacao,
        # e o construtor nao pode inventar.
        try:
            construir_conjunto(roteiro, d, "perfil", modos=("padrao",),
                               perfis=("libras",))
            print("  %-28s ERRO: aceitou perfil sem mensagem_chave" % "perfil sem chave")
            falhas += 1
        except ErroDeRoteiro as e:
            ok = "mensagem_chave" in str(e)
            print("  %-28s %s" % ("perfil sem chave",
                                  "recusou corretamente" if ok else "ERRO: %s" % e))
            falhas += 0 if ok else 1

        com_chave = copy_roteiro_com_chave(roteiro)
        perfis = construir_conjunto(com_chave, d, "perfis", modos=("padrao",),
                                    perfis=("completo", "libras", "leitura_facil"))
        for chave, caminho in perfis.items():
            falhas += len(auditar(caminho, "deck %s" % chave))
            falhas += conta_slides(caminho, esperado, "deck %s" % chave)
        falhas += testa_recursos_por_perfil(perfis, d)
        falhas += testa_paridade(list(feitos.values()), d)

        unico = construir_conjunto(roteiro, d, "legado", arquivo_unico=True)
        alvo = unico["arquivo_unico"]
        falhas += len(auditar(alvo, "deck legado (arquivo unico)"))
        falhas += conta_slides(alvo, 1 + 3 * esperado, "deck legado")

    print("\nO construtor recusa roteiro que produziria slide inacessivel:")
    import copy
    checagens = []

    r = copy.deepcopy(roteiro)
    r["slides"][2]["titulo"] = r["slides"][1]["titulo"]
    checagens.append(testa_recusa(r, "titulo repetido"))

    r = copy.deepcopy(roteiro)
    r["slides"][2]["conteudo"] = ["linha"] * 7
    checagens.append(testa_recusa(r, "limite e 6"))

    r = copy.deepcopy(roteiro)
    r["slides"][4]["links"] = [{"texto": "clique aqui", "url": "https://exemplo.org"}]
    checagens.append(testa_recusa(r, "link nao descritivo"))

    r = copy.deepcopy(roteiro)
    r["apresentacao"]["titulo"] = ""
    checagens.append(testa_recusa(r, "apresentacao.titulo"))

    falhas += sum(1 for c in checagens if not c)

    print("\n%s" % ("TUDO OK" if not falhas else "%d PROBLEMA(S)" % falhas))
    return 1 if falhas else 0


if __name__ == "__main__":
    sys.exit(main())
