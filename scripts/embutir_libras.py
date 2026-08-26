"""
embutir_libras — poe a janela de Libras DENTRO do deck.

Mesma razao do `embutir_audio.py`: entregar o vídeo numa pasta ao lado cumpre a
regra J01, e quase ninguem abre a pasta. A pergunta "como acesso a Libras?" e a
prova de que o recurso existia e nao estava ao alcance.

O video vai na CAPA, na faixa livre a direita: e o primeiro slide que qualquer
pessoa ve, e a capa e o unico layout com meia largura sobrando — em qualquer
outro ele cobriria texto e reprovaria em N03.

Cuidados que vem do catalogo:
  I04  sem autoplay e sem loop — quem decide assistir e o usuario
  D10  o objeto de midia recebe texto alternativo
  J05  a glosa e automatica: o alt text DIZ isso, nao esconde
  N01  proporcao preservada; a janela nao pode ser esticada
  N02  dentro da area segura, fora da faixa do rodape
  O02  vai nos TRES modos, senao uma versao fica sem o recurso

Precisa do PowerPoint instalado (COM), porque o python-pptx nao insere video.

    python scripts/embutir_libras.py deck.pptx libras/janela-libras.mp4
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import grade as G

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

PT = 12700                       # COM trabalha em pontos: 1 pt = 12700 EMU
NATIVO = (312, 452)              # o recorte que libras_caminho_a.py produz

ALT = ("Janela de Libras com o resumo deste material. A glosa é automática "
       "e não foi revisada por intérprete; o roteiro em SRT acompanha.")


def _caixa(nativo, faixa_cheia=False):
    """
    Onde a janela entra. Na capa ela usa 10 cm de altura; no perfil de Libras,
    a faixa reservada inteira — e ai ela passa nos dois minimos da NBR 15290
    por analogia (8,47 cm de largura e 9,53 cm de altura), com 9,11 x 13,20.
    """
    caixa_w = G.larg(4)
    caixa_h = G.CONTEUDO_H if faixa_cheia else G.cm(10.0)
    w, h, dx, dy = G.encaixar(nativo[0], nativo[1], caixa_w, caixa_h)
    x = G.x(8) + dx
    y = G.CONTEUDO_Y + dy
    return x, y, w, h


def _tamanho_nativo(video):
    """Le a resolucao real do arquivo; sem ffprobe, usa o recorte conhecido."""
    import shutil
    import subprocess

    if not shutil.which("ffprobe"):
        return NATIVO
    try:
        saida = subprocess.run(
            ["ffprobe", "-v", "error", "-select_streams", "v:0",
             "-show_entries", "stream=width,height", "-of", "csv=p=0",
             video], capture_output=True, text=True, timeout=30).stdout
        w, h = (int(v) for v in saida.strip().split(",")[:2])
        return (w, h) if w and h else NATIVO
    except Exception:
        return NATIVO


def embutir(pptx: str, video: str, slide_no: int = 1) -> dict:
    import win32com.client as win32

    pptx, video = os.path.abspath(pptx), os.path.abspath(video)
    if not os.path.exists(video):
        raise RuntimeError("video nao encontrado: %s" % video)

    x, y, w, h = _caixa(_tamanho_nativo(video))

    app = win32.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(pptx, WithWindow=False)
    try:
        slide = pres.Slides(slide_no)
        # se ja houver uma janela embutida, troca em vez de empilhar
        for forma in list(slide.Shapes):
            if forma.Name == "Janela de Libras":
                forma.Delete()
        midia = slide.Shapes.AddMediaObject2(
            video, LinkToFile=0, SaveWithDocument=-1,
            Left=x // PT, Top=y // PT, Width=w // PT, Height=h // PT)
        midia.Name = "Janela de Libras"
        midia.AlternativeText = ALT
        try:
            fmt = midia.AnimationSettings.PlaySettings
            fmt.PlayOnEntry = 0          # sem autoplay (regra I04)
            fmt.LoopUntilStopped = 0
            fmt.HideWhileNotPlaying = 0
        except Exception:
            pass
        pres.Save()
    finally:
        try:
            pres.Close()
        except Exception:
            pass

    return {"slide": slide_no, "bytes": os.path.getsize(pptx),
            "caixa_cm": (round(w / G.cm(1), 1), round(h / G.cm(1), 1))}


def embutir_por_slide(pptx: str, pasta: str) -> dict:
    """
    Uma janela POR SLIDE, a partir do manifesto de `gen_libras_slides.py`.

    A ordem no spTree importa: o audio ja quebrou C01, C02 e N04 uma vez por
    entrar no fim da arvore, e aqui sao 28 objetos, nao um. Cada janela vai
    para logo depois do titulo — que e o lugar semantico certo, "a Libras
    DESTE slide".
    """
    import json

    import win32com.client as win32

    pptx = os.path.abspath(pptx)
    caminho_man = os.path.join(os.path.abspath(pasta), "manifesto.json")
    if not os.path.exists(caminho_man):
        raise RuntimeError("manifesto ausente: rode gen_libras_slides.py antes")
    with open(caminho_man, encoding="utf-8") as f:
        manifesto = json.load(f)

    por_slide = {}
    for m in manifesto.values():
        if m.get("arquivo") and os.path.exists(m["arquivo"]):
            por_slide[m["slide"]] = os.path.abspath(m["arquivo"])
    if not por_slide:
        raise RuntimeError("nenhum video no manifesto")

    nativo = _tamanho_nativo(list(por_slide.values())[0])
    x, y, w, h = _caixa(nativo, faixa_cheia=True)

    app = win32.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(pptx, WithWindow=False)
    postos = 0
    try:
        for i in range(1, pres.Slides.Count + 1):
            video = por_slide.get(i)
            if not video:
                continue
            slide = pres.Slides(i)
            for forma in list(slide.Shapes):
                if forma.Name == "Janela de Libras":
                    forma.Delete()
            midia = slide.Shapes.AddMediaObject2(
                video, LinkToFile=0, SaveWithDocument=-1,
                Left=x // PT, Top=y // PT, Width=w // PT, Height=h // PT)
            midia.Name = "Janela de Libras"
            midia.AlternativeText = ALT
            try:
                fmt = midia.AnimationSettings.PlaySettings
                fmt.PlayOnEntry = 0
                fmt.LoopUntilStopped = 0
                fmt.HideWhileNotPlaying = 0
            except Exception:
                pass
            postos += 1
        pres.Save()
    finally:
        try:
            pres.Close()
        except Exception:
            pass

    _reordenar(pptx)
    return {"slides_com_janela": postos, "bytes": os.path.getsize(pptx),
            "caixa_cm": (round(w / G.cm(1), 1), round(h / G.cm(1), 1))}


def _reordenar(pptx: str) -> int:
    """Poe a janela logo DEPOIS do titulo no spTree (regras C01 e N04)."""
    from pptx import Presentation

    import a11y_lib as A

    prs = Presentation(pptx)
    movidos = 0
    for slide in prs.slides:
        spTree = slide.shapes._spTree
        titulo = janela = None
        for el in A.iter_shape_elements(spTree):
            if A.is_title_placeholder(el):
                titulo = el
            elif A.media_kind(el) == "video":
                janela = el
        if titulo is None or janela is None:
            continue
        filhos = list(spTree)
        if filhos.index(janela) == filhos.index(titulo) + 1:
            continue
        spTree.remove(janela)
        spTree.insert(list(spTree).index(titulo) + 1, janela)
        movidos += 1
    prs.save(pptx)
    return movidos


def main():
    ap = argparse.ArgumentParser(description="Embute a janela de Libras no deck")
    ap.add_argument("pptx")
    ap.add_argument("video")
    ap.add_argument("--slide", type=int, default=1)
    ap.add_argument("--por-slide", action="store_true",
                    help="`video` e a PASTA com o manifesto: uma janela por slide")
    args = ap.parse_args()

    try:
        if args.por_slide:
            r = embutir_por_slide(args.pptx, args.video)
            print("janela de Libras em %d slides · %.1f × %.1f cm · %.1f MB"
                  % (r["slides_com_janela"], r["caixa_cm"][0],
                     r["caixa_cm"][1], r["bytes"] / 1048576))
            return 0
        r = embutir(args.pptx, args.video, args.slide)
    except Exception as e:
        print("ERRO: %s" % e)
        print("\nExige Windows com PowerPoint instalado. Sem ele, entregue o")
        print("vídeo na pasta ao lado e diga no LEIA-ME onde ele está.")
        return 2

    print("janela de Libras no slide %d · %.1f × %.1f cm · arquivo com %.1f MB"
          % (r["slide"], r["caixa_cm"][0], r["caixa_cm"][1],
             r["bytes"] / 1048576))
    print("\nSem reprodução automática. A glosa é automática e o alt text diz "
          "isso (regra J05).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
