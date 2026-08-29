"""
audit_pptx — auditor estatico de apresentacoes acessiveis.

Implementa as camadas A a I do catalogo (references/02-catalogo-auditoria.md).
As camadas J a M dependem de midia, do PDF e de humano, e saem no relatorio
como NAO VERIFICADO — nunca como aprovadas.

Uso:
    python audit_pptx.py deck.pptx [--json saida.json] [--md relatorio.md]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import a11y_lib as A
import audit_contrast as C
import audit_design

# WCAG 1.4.8 (AAA): "Width is no more than 80 characters or glyphs".
# O codigo usava 90 e a mensagem dizia ~70 — uma folga de 20 caracteres que nao
# estava documentada em lugar nenhum. 80 tem fonte; 70 e 90 nao tinham.
LIMITE_LINHA = 80

EMU_MIN_TARGET = 228600  # 24 px CSS a 96 dpi = 0,25 pol

CONFORME = "conforme"
NAO_CONFORME = "nao_conforme"
NAO_VERIFICADO = "nao_verificado"


class Report:
    def __init__(self, path):
        self.path = path
        self.findings = []
        self.evaluated = set()

    def check(self, rule):
        self.evaluated.add(rule)

    def fail(self, rule, sev, criterio, onde, detalhe):
        self.evaluated.add(rule)
        self.findings.append({
            "regra": rule, "severidade": sev, "criterio": criterio,
            "veredito": NAO_CONFORME, "onde": onde, "detalhe": detalhe,
        })

    def unverified(self, rule, sev, criterio, detalhe, onde="—"):
        self.findings.append({
            "regra": rule, "severidade": sev, "criterio": criterio,
            "veredito": NAO_VERIFICADO, "onde": onde, "detalhe": detalhe,
        })

    def counts(self):
        c = defaultdict(int)
        for f in self.findings:
            if f["veredito"] == NAO_CONFORME:
                c[f["severidade"]] += 1
            else:
                c["nao_verificado"] += 1
        return dict(c)


def onde(slide_no, el=None, extra=""):
    s = "slide %d" % slide_no
    if el is not None:
        nome = A.shape_name(el) or A.local(el)
        s += " · %s" % nome
    if extra:
        s += " · %s" % extra
    return s


# ==========================================================================
# Camada A — documento e metadados
# ==========================================================================
def audit_metadata(prs, rep: Report, caminho: str):
    cp = prs.core_properties
    base = os.path.basename(caminho)
    stem = os.path.splitext(base)[0]

    rep.check("A01")
    titulo = (cp.title or "").strip()
    if not titulo:
        rep.fail("A01", "E", "2.4.2", base, "dc:title vazio — o leitor de tela "
                 "anuncia o nome do arquivo, e o PDF exportado herda o defeito (L03)")
    elif titulo.lower() == stem.lower():
        rep.fail("A01", "E", "2.4.2", base, "dc:title igual ao nome do arquivo: %r" % titulo)

    rep.check("A02")
    lang = (cp.language or "").strip()
    if not lang:
        rep.fail("A02", "E", "3.1.1", base, "idioma do documento nao definido")
    elif not lang.lower().startswith("pt"):
        rep.fail("A02", "E", "3.1.1", base, "idioma do documento e %r" % lang)

    rep.check("A04")
    autor = (cp.author or "").strip()
    if not autor or autor.lower() in ("usuario do windows", "usuário do windows",
                                      "windows user", "user", "autor"):
        rep.fail("A04", "A", "—", base, "autor ausente ou generico: %r" % autor)

    rep.check("A05")
    if re.match(r"^(apresenta[cç][aã]o|presentation|slide|deck|documento)\s*\d*$",
                stem.strip(), re.I):
        rep.fail("A05", "A", "—", base, "nome de arquivo nao descritivo")

    rep.check("A06")
    w, h = A.slide_size(prs)
    if h and abs((w / h) - (16 / 9)) > 0.02:
        rep.fail("A06", "D", "—", base,
                 "proporcao %.2f:1, fora de 16:9" % (w / h))


def audit_run_language(prs, rep: Report):
    rep.check("A03")
    ruins = defaultdict(list)
    for i, slide in enumerate(prs.slides, 1):
        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            for p_el, _cel in A.iter_all_paragraphs(el):
                for r_el in A.iter_runs(p_el):
                    txt = A.run_text(r_el).strip()
                    if not txt:
                        continue
                    lang = A.run_lang(r_el)
                    if lang is None:
                        ruins["ausente"].append((i, txt[:40]))
                    elif not lang.lower().startswith("pt"):
                        ruins[lang].append((i, txt[:40]))
    for lang, itens in ruins.items():
        slides = sorted({s for s, _ in itens})
        rep.fail("A03", "E", "3.1.2",
                 "slides %s" % ", ".join(map(str, slides[:12])),
                 "%d trechos com lang=%s (ex.: %r) — sintese de voz sai com "
                 "fonetica errada" % (len(itens), lang, itens[0][1]))


# ==========================================================================
# Camada B — estrutura semantica
# ==========================================================================
def audit_structure(prs, rep: Report):
    for r in ("B01", "B02", "B03", "B05", "B06", "B09"):
        rep.check(r)
    titulos = {}
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(A.iter_shape_elements(slide.shapes._spTree))
        title_el = next((s for s in shapes if A.is_title_placeholder(s)), None)

        if title_el is None:
            rep.fail("B01", "E", "1.3.1 · 2.4.2", onde(i),
                     "slide sem placeholder de titulo — nao aparece no indice "
                     "de titulos do leitor de tela")
        else:
            txt = " ".join(A.paragraph_text(p) for p in A.iter_paragraphs(title_el)).strip()
            if not txt:
                rep.fail("B02", "E", "2.4.2", onde(i, title_el), "titulo vazio")
            else:
                titulos.setdefault(txt.lower(), []).append(i)

        layout_nome = ""
        try:
            layout_nome = slide.slide_layout.name or ""
        except Exception:
            pass
        if re.search(r"^(em branco|blank)$", layout_nome.strip(), re.I):
            rep.fail("B06", "A", "1.3.1", onde(i),
                     "layout 'Em Branco' — sem placeholders, sem semantica")

        for el in shapes:
            if A.local(el) == "sp" and A.get_ph(el) is None:
                txt = " ".join(A.paragraph_text(p) for p in A.iter_paragraphs(el)).strip()
                if txt:
                    rep.fail("B05", "E", "1.3.1", onde(i, el),
                             "texto em caixa solta, fora de placeholder: %r" % txt[:50])

        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            for p_el in A.iter_paragraphs(el):
                txt = A.paragraph_text(p_el).strip()
                pp = A.paragraph_props(p_el)
                if re.match(r"^[-*•·]\s+\S", txt) and pp["bullet"] in (None, "buNone"):
                    rep.fail("B09", "A", "1.3.1", onde(i, el),
                             "lista digitada com hifen/asterisco em vez de "
                             "marcador nativo: %r" % txt[:40])

    for txt, slides in titulos.items():
        if len(slides) > 1:
            rep.fail("B03", "E", "—", "slides %s" % ", ".join(map(str, slides)),
                     "titulo repetido %r — use sufixo de continuidade, "
                     "como '(1 de %d)'" % (txt[:50], len(slides)))


def audit_sections(prs, rep: Report):
    rep.check("B07")
    rep.check("B08")
    secoes = A.get_sections(prs)
    if not secoes:
        return
    vistos = defaultdict(int)
    for nome, _ in secoes:
        n = (nome or "").strip()
        if not n or re.match(r"^(se[cç][aã]o|section)\s*(padr[aã]o|default)?\s*\d*$", n, re.I):
            rep.fail("B07", "E", "—", "secao %r" % n, "nome de secao sem significado")
        vistos[n.lower()] += 1
    for n, c in vistos.items():
        if c > 1:
            rep.fail("B08", "D", "—", "secao %r" % n, "nome de secao repetido %d vezes" % c)


# ==========================================================================
# Camada C — ordem de leitura
# ==========================================================================
def audit_reading_order(prs, rep: Report):
    for r in ("C01", "C02", "C04", "C05", "C06"):
        rep.check(r)
    slide_w, slide_h = A.slide_size(prs)
    banda = max(1, slide_h // 12)

    for i, slide in enumerate(prs.slides, 1):
        shapes = list(A.iter_shape_elements(slide.shapes._spTree))
        relevantes = [s for s in shapes
                      if not A.is_decorative(s) and not A.is_hidden(s)]

        if relevantes:
            primeiro = relevantes[0]
            if not A.is_title_placeholder(primeiro):
                tem_titulo = any(A.is_title_placeholder(s) for s in relevantes)
                if tem_titulo:
                    rep.fail("C01", "E", "1.3.2", onde(i, primeiro),
                             "o titulo nao e o primeiro na ordem de leitura; "
                             "o leitor anuncia %r antes" % (A.shape_name(primeiro) or "?"))

        # posicao RESOLVIDA (slide -> layout -> master): ler so o xfrm do
        # slide deixava passar o slide cujos placeholders herdam a posicao
        posicionados = [(s, A.resolve_xfrm(slide, s)[0]) for s in relevantes]
        posicionados = [(s, b) for s, b in posicionados if b is not None]
        # Objeto de midia e CROMO da interface, nao conteudo da coluna: um
        # controle de audio no canto nao deve invalidar a leitura por coluna
        # de uma comparacao lado a lado.
        posicionados = [(s, b) for s, b in posicionados if not A.media_kind(s)]
        if len(posicionados) > 1:
            atual = [s for s, _ in posicionados]
            esperado = [s for s, _ in sorted(
                posicionados, key=lambda t: (t[1][1] // banda, t[1][0]))]
            por_coluna = [s for s, _ in sorted(
                posicionados, key=lambda t: (t[1][0], t[1][1]))]
            colunas_paralelas = len({b[0] for _s, b in posicionados}) >= 2
            if atual != esperado and not (atual == por_coluna
                                          and colunas_paralelas):
                fora = [A.shape_name(s) or A.local(s) for s in atual[:6]]
                rep.fail("C02", "E", "1.3.2", onde(i),
                         "ordem do spTree diverge do fluxo visual "
                         "(cima→baixo, esquerda→direita). Ordem atual: %s" % ", ".join(fora))

        for el in shapes:
            n = A.count_group_children(el)
            if n > 5 and not A.get_alt_text(el).strip() and not A.is_decorative(el):
                rep.fail("C04", "A", "1.3.2", onde(i, el),
                         "grupo com %d objetos sem alt text no grupo — o leitor "
                         "anuncia cada peca isoladamente" % n)
            nome = A.shape_name(el)
            if re.match(r"^(rectangle|oval|picture|imagem|text ?box|caixa de texto|"
                        r"retangulo|retângulo|elipse|forma|shape|group|grupo|"
                        r"content placeholder|espa[cç]o reservado)\s*\d+$", nome.strip(), re.I):
                rep.fail("C05", "A", "—", onde(i, el),
                         "nome automatico %r — dificulta a navegacao no Painel de Selecao" % nome)
            box_res = A.resolve_xfrm(slide, el)[0]
            fora = (box_res is not None and (box_res[0] + box_res[2] <= 0
                    or box_res[1] + box_res[3] <= 0
                    or box_res[0] >= slide_w or box_res[1] >= slide_h))
            if fora and not A.is_decorative(el):
                if not A.is_title_placeholder(el):
                    rep.fail("C06", "A", "1.3.2", onde(i, el),
                             "objeto fora da area do slide mas ainda na ordem de leitura")


# ==========================================================================
# Camada D — texto alternativo
# ==========================================================================
def needs_alt(el) -> bool:
    t = A.local(el)
    if t in ("pic", "graphicFrame", "grpSp", "cxnSp"):
        return True
    if t == "sp":
        if A.get_ph(el) is not None:
            return False
        texto = " ".join(A.paragraph_text(p) for p in A.iter_paragraphs(el)).strip()
        return not texto
    return False


def audit_alt_text(prs, rep: Report):
    for r in ("D01", "D02", "D03", "D04", "D05", "D06", "D08", "D10", "D11"):
        rep.check(r)
    for i, slide in enumerate(prs.slides, 1):
        notas = ""
        try:
            if slide.has_notes_slide:
                notas = slide.notes_slide.notes_text_frame.text or ""
        except Exception:
            pass

        for el in A.iter_shape_elements(slide.shapes._spTree):
            if not needs_alt(el):
                continue
            alt = A.get_alt_text(el)
            dec = A.is_decorative(el)

            if dec and alt.strip():
                rep.fail("D11", "E", "1.1.1", onde(i, el),
                         "marcado como decorativo E com alt text %r — o objeto diz "
                         "'me ignore' e 'me leia' ao mesmo tempo" % alt[:40])
                continue
            if dec:
                continue

            if not alt:
                rep.fail("D01", "E", "1.1.1", onde(i, el),
                         "sem texto alternativo e sem marca de decorativo")
                continue
            if not alt.strip():
                rep.fail("D02", "E", "1.1.1", onde(i, el), "alt text so com espacos")
                continue
            if A.looks_like_filename(alt) or A.is_generic_alt(alt):
                rep.fail("D03", "E", "1.1.1", onde(i, el),
                         "alt text e nome de arquivo ou rotulo generico: %r" % alt[:60])
            if A.RE_AI_RESIDUE.search(alt):
                rep.fail("D05", "E", "1.1.1", onde(i, el),
                         "residuo de descricao automatica: %r" % alt[:80])
            if A.RE_ALT_PREFIX.search(alt):
                rep.fail("D04", "A", "1.1.1", onde(i, el),
                         "comeca com o tipo do objeto (%r) — o leitor de tela ja "
                         "anuncia o tipo" % alt[:40])
            if len(alt) > 150:
                rep.fail("D06", "A", "1.1.1", onde(i, el),
                         "alt text com %d caracteres; acima de 150 alguns leitores "
                         "truncam — leve o detalhe para as Notas" % len(alt))
            # Nao basta "nao tem acento": frase legitima pode nao ter nenhuma
            # palavra acentuada. Acusa so quando ha palavra que EXIGE acento.
            if A.falta_acento(alt):
                rep.fail("D08", "A", "3.1.1", onde(i, el),
                         "alt text com palavra que deveria estar acentuada: %r"
                         % alt[:70])

            kind = A.media_kind(el)
            if kind and not alt.strip():
                rep.fail("D10", "A", "1.1.1", onde(i, el), "%s sem descricao" % kind)

            # D07 vale para grafico/diagrama, nao para tabela: a tabela ja e
            # acessivel por estrutura propria (camada G) e nao precisa de
            # descricao longa.
            if _e_figura_densa(el) and not notas.strip():
                rep.unverified("D07", "E", "1.1.1",
                               "figura densa sem descricao longa nas Notas do slide",
                               onde(i, el))


_URI_CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_URI_DIAGRAM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def _e_figura_densa(el) -> bool:
    """Grafico ou SmartArt — nao tabela, nao imagem simples."""
    if A.local(el) != "graphicFrame":
        return False
    if A.get_tables(el):
        return False
    for gd in el.iter(A.q("a:graphicData")):
        if gd.get("uri") in (_URI_CHART, _URI_DIAGRAM):
            return True
    return False


_PT_HINTS = re.compile(r"\b(de|da|do|com|para|que|uma?|os|as|no|na|em|e|dos|das)\b", re.I)


def _parece_pt(s: str) -> bool:
    return len(_PT_HINTS.findall(s)) >= 2


# ==========================================================================
# Camada E — cor e contraste
# ==========================================================================
def audit_contrast_rules(prs, rep: Report):
    for r in ("E01", "E02", "E03", "E08"):
        rep.check(r)
    resolver = C.ThemeResolver(prs)

    for i, slide in enumerate(prs.slides, 1):
        bg, bg_desc = C.slide_background(slide, resolver)

        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            if A.is_decorative(el):
                continue
            fill, fill_desc = C.shape_fill(el, slide, resolver)
            fundo, fundo_desc = (fill, fill_desc) if fill else (bg, bg_desc)

            # celula de tabela tem o proprio fundo: usar o dela, nao o do slide
            for p_el, cel in A.iter_all_paragraphs(el):
                if cel is not None:
                    fundo, fundo_desc = C.cell_fill(cel["tc"], slide, resolver)
                else:
                    fundo, fundo_desc = (fill, fill_desc) if fill else (bg, bg_desc)
                for r_el in A.iter_runs(p_el):
                    txt = A.run_text(r_el).strip()
                    if not txt:
                        continue
                    sz, _origem = A.resolve_font_size(slide, el, p_el, r_el)
                    bold = A.resolve_bold(el, p_el, r_el)
                    cor, cor_desc = C.run_color(r_el, slide, resolver)
                    if cor is None:
                        cor, cor_desc = C.default_text_color(slide, resolver)

                    if cor is None or fundo is None:
                        motivo = fundo_desc if fundo is None else cor_desc
                        rep.unverified(
                            "E09" if motivo in ("gradFill", "blipFill", "pattFill")
                            or "blipFill" in str(motivo) or "gradFill" in str(motivo)
                            else "E01",
                            "A", "1.4.3",
                            "contraste nao calculavel automaticamente (%s) — texto %r "
                            "exige conferencia manual com conta-gotas" % (motivo, txt[:30]),
                            onde(i, el))
                        continue

                    ratio = C.contrast_ratio(cor, fundo)
                    grande = C.is_large_text(sz, bold)
                    exigido = C.required_ratio(sz, bold, "AA")
                    if ratio + 0.005 < exigido:
                        rep.fail("E02" if grande else "E01", "E", "1.4.3", onde(i, el),
                                 "contraste %.2f:1 (exigido %.1f:1) — texto %s %r, "
                                 "%s sobre %s" % (
                                     ratio, exigido,
                                     "grande" if grande else "normal", txt[:30],
                                     C.rgb_to_hex(cor), C.rgb_to_hex(fundo)))
                    elif not grande and ratio < 7.0:
                        rep.fail("E03", "D", "1.4.6", onde(i, el),
                                 "contraste %.2f:1 — abaixo da meta AAA de 7:1" % ratio)

                    if C.rgb_to_hex(cor) == "#000000" and C.rgb_to_hex(fundo) == "#FFFFFF":
                        rep.fail("E08", "A", "—", onde(i, el),
                                 "preto puro sobre branco puro — irradiacao e fadiga "
                                 "visual; prefira #1A1A1A sobre #FAF7F2")


# ==========================================================================
# Camada F — tipografia
# ==========================================================================
# F10 — sigla expandida na primeira ocorrencia (WCAG 3.1.4, AAA)
#
# Ate 26/08/2026 esta regra estava em rep.check e NAO tinha nenhum rep.fail:
# saia do relatorio como CONFORME sem nunca ter sido verificada. E o defeito
# que este projeto existe para recusar, e estava dentro do proprio auditor.
#
# Conta como expansao: "Nome Por Extenso (SIGLA)", "SIGLA (Nome Por Extenso)",
# ou uma linha de glossario nas notas do orador ("SIGLA: nome por extenso").
# A busca varre o deck INTEIRO, notas incluidas, porque a expansao pode estar
# na nota do slide em que a sigla aparece.
# ==========================================================================
RE_SIGLA = re.compile(r"\b([A-ZÀ-Ý][A-ZÀ-Ý0-9]{2,7})\b")

# Siglas que sao o nome corrente da coisa: expandi-las atrapalha em vez de
# ajudar. Nao e conveniencia — e o proprio criterio 3.1.4, que fala de
# "abbreviations", nao de nomes proprios consagrados.
SIGLAS_DISPENSADAS = {
    "PDF", "HTML", "XML", "CSS", "PNG", "JPG", "GIF", "MP3", "MP4", "SRT",
    "YAML", "JSON", "ZIP", "URL", "API", "CLI", "GPU", "CPU", "USB", "USP",
    "UFF", "ABNT", "NBR", "ISO", "WCAG", "PPTX", "DOCX", "OOXML", "EMU",
    "AAA", "III", "PT", "BR", "EUA", "W3C", "IFLA", "ITU", "CNMP",
}


# Nem todo bloco de maiusculas e sigla. Estes tres casos apareceram no proprio
# deck do projeto e sao ruido, nao achado:
#   NAO, SIM, ...   palavra portuguesa em caixa alta por enfase
#   D05, K03, ...   identificador de regra do proprio catalogo
#   e-MAG           a sigla ja vem colada a um prefixo minusculo
PALAVRAS_CAIXA_ALTA = {"NÃO", "NAO", "SIM", "TODOS", "NUNCA", "SEMPRE", "MAS",
                       "OU", "SEM", "COM", "POR", "PARA", "UMA", "DOIS", "TRES"}


def _sigla_dispensada(sigla, contexto=""):
    if sigla in SIGLAS_DISPENSADAS or sigla in PALAVRAS_CAIXA_ALTA:
        return True
    # identificador de regra: uma letra e dois digitos (D05, K03, N11)
    if len(sigla) == 3 and sigla[0].isalpha() and sigla[1:].isdigit():
        return True
    # sigla colada a prefixo minusculo, como o "MAG" de "e-MAG"
    if re.search(r"[a-z]-%s" % sigla, contexto):
        return True
    # nome de arquivo, como o "SKILL" de "SKILL.md"
    if re.search(r"%s\.[a-z]{2,4}\b" % sigla, contexto):
        return True
    return False



def _expansoes_do_deck(prs):
    """Siglas que aparecem expandidas em algum lugar do material."""
    textos = []
    for slide in prs.slides:
        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            for p_el, _ in A.iter_all_paragraphs(el):
                t = A.paragraph_text(p_el).strip()
                if t:
                    textos.append(t)
        try:
            nota = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            nota = ""
        if nota:
            textos.append(nota)

    achadas = set()
    inteiro = " ".join(textos)
    for sigla in set(RE_SIGLA.findall(inteiro)):
        padroes = (
            r"\(%s\)" % sigla,                       # Nome Por Extenso (SIGLA)
            r"%s\s*\([^)]{4,}\)" % sigla,             # SIGLA (Nome Por Extenso)
            r"%s\s*[:—-]\s*\w{4,}" % sigla,           # glossario: SIGLA — nome
        )
        if any(re.search(pad, inteiro) for pad in padroes):
            achadas.add(sigla)
    return achadas


# ==========================================================================
def audit_typography(prs, rep: Report):
    for r in ("F01", "F02", "F03", "F04", "F05", "F06", "F07", "F10"):
        rep.check(r)
    siglas_vistas = set()
    expandidas = _expansoes_do_deck(prs)

    for i, slide in enumerate(prs.slides, 1):
        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            titulo = A.is_title_placeholder(el)
            paras_forma = list(A.iter_paragraphs(el))
            n_bullets = sum(1 for p in paras_forma if A.paragraph_text(p).strip())
            # celula de tabela tambem e texto: a 15pt passava batido antes
            paras = [(p, None) for p in paras_forma]
            paras += [(p, cel) for p, cel in A.iter_all_paragraphs(el)
                      if cel is not None]

            if not titulo and n_bullets > 6:
                rep.fail("F07", "A", "—", onde(i, el),
                         "%d paragrafos num so bloco — carga cognitiva; divida o slide"
                         % n_bullets)

            for p_el, celula in paras:
                txt = A.paragraph_text(p_el).strip()
                if not txt:
                    continue
                pp = A.paragraph_props(p_el)
                onde_txt = onde(i, el, "célula %d,%d" % (celula["ri"], celula["ci"])
                                if celula else "")

                if pp["algn"] == "just":
                    rep.fail("F04", "E", "—", onde_txt,
                             "paragrafo justificado — cria 'rios de branco' que "
                             "interrompem a leitura: %r" % txt[:40])
                # F05 vale para CORPO de texto. Entrelinha apertada em titulo de
                # display e tipografia normal e nao produz o efeito de troca de
                # linha que a regra combate; exigir 1,5 num titulo de duas linhas
                # so afasta as linhas sem ganho de legibilidade.
                # Nem titulo nem celula de tabela: a entrelinha de 1,5 combate
                # a troca involuntaria de linha ao ler PARAGRAFO. Numa celula com
                # um valor curto, ela so infla a linha e atrapalha a varredura.
                # tamanho do primeiro run, para saber se e display ou leitura
                _szs = [A.resolve_font_size(slide, el, p_el, r)[0]
                        for r in A.iter_runs(p_el)]
                _szs = [z for z in _szs if z]
                display = bool(_szs) and max(_szs) >= 4000
                if (not titulo and celula is None and not display
                        and pp["line_pct"] is not None
                        and pp["line_pct"] < 150000):
                    rep.fail("F05", "A", "—", onde(i, el),
                             "entrelinha %.2f no corpo de texto, abaixo de 1,5"
                             % (pp["line_pct"] / 100000))
                if titulo and pp["line_pct"] is not None and pp["line_pct"] < 90000:
                    rep.fail("F05", "A", "—", onde(i, el),
                             "entrelinha %.2f no titulo — abaixo de 0,9 as linhas colidem"
                             % (pp["line_pct"] / 100000))
                if len(txt) > LIMITE_LINHA and celula is None:
                    rep.fail("F07", "A", "1.4.8", onde_txt,
                             "linha com %d caracteres — o WCAG 1.4.8 pede no "
                             "máximo %d; acima disso o olho perde o retorno de "
                             "linha" % (len(txt), LIMITE_LINHA))
                cap_herdado, cap_origem = A.resolve_caps(slide, el, p_el)
                if cap_herdado in ("all", "small"):
                    rep.fail("F06", "A", "—", onde_txt,
                             "renderizado em CAIXA ALTA por heranca do %s "
                             "(cap=%r): %r" % (cap_origem, cap_herdado, txt[:40]))
                if A.is_all_caps_sentence(txt):
                    rep.fail("F06", "A", "—", onde_txt,
                             "frase inteira em CAIXA ALTA: %r — suprime ascendentes e "
                             "descendentes" % txt[:40])
                for sigla in RE_SIGLA.findall(txt):
                    if sigla in siglas_vistas:
                        continue
                    siglas_vistas.add(sigla)
                    if _sigla_dispensada(sigla, txt) or sigla in expandidas:
                        continue
                    rep.fail("F10", "A", "3.1.4", onde_txt,
                             "sigla %r usada sem ser expandida em lugar nenhum "
                             "do material — expanda na primeira ocorrência ou "
                             "ponha um glossário nas notas" % sigla)

                for r_el in A.iter_runs(p_el):
                    rtxt = A.run_text(r_el).strip()
                    if not rtxt:
                        continue
                    props = A.run_props(r_el)
                    face = props["typeface"]
                    if face and face.strip().lower().lstrip("+") not in A.SANS_SAFE:
                        if not face.startswith("+"):
                            rep.fail("F01", "A", "—", onde(i, el),
                                     "fonte %r fora do conjunto seguro sem serifa" % face)
                    sz, origem = A.resolve_font_size(slide, el, p_el, r_el)
                    if sz is None:
                        rep.unverified("F02", "E", "1.4.4",
                                       "tamanho de fonte indeterminado (%s) para %r"
                                       % (origem, rtxt[:30]), onde_txt)
                    elif titulo and sz < 3200:
                        rep.fail("F03", "A", "—", onde_txt,
                                 "titulo com %.0fpt, abaixo de 32pt" % (sz / 100))
                    elif not titulo and sz < 1800:
                        rep.fail("F02", "E", "1.4.4", onde_txt,
                                 "corpo com %.0fpt, abaixo do minimo de 18pt "
                                 "(fonte: %s)" % (sz / 100, origem))
                    if props["i"]:
                        rep.fail("F06", "A", "—", onde(i, el),
                                 "italico em %r — fragmenta as hastes das letras" % rtxt[:30])
                    if props["u"] and props["u"] != "none":
                        rid, _ = A.get_hyperlink(el)
                        rPr = props["rPr"]
                        tem_link = rPr is not None and rPr.find(A.q("a:hlinkClick")) is not None
                        if not tem_link and not rid:
                            rep.fail("F06", "A", "—", onde(i, el),
                                     "sublinhado fora de hiperlink em %r" % rtxt[:30])


# ==========================================================================
# Camada G — tabelas
# ==========================================================================
def audit_tables(prs, rep: Report):
    for r in ("G01", "G03", "G04", "G05"):
        rep.check(r)
    for i, slide in enumerate(prs.slides, 1):
        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            for tbl in A.get_tables(el):
                tp = A.table_props(tbl)
                if not tp["first_row"]:
                    rep.fail("G01", "E", "1.3.1", onde(i, el),
                             "tabela %dx%d sem linha de cabecalho (firstRow) — o leitor "
                             "de tela nao consegue recontextualizar as celulas"
                             % (tp["n_rows"], tp["n_cols"]))
                if tp["merged"]:
                    rep.fail("G03", "E", "1.3.1", onde(i, el),
                             "%d celulas mescladas/divididas — corrompe a contagem de "
                             "colunas do leitor de tela" % len(tp["merged"]))
                if tp["empty_rows"] or tp["empty_cols"]:
                    rep.fail("G04", "A", "1.3.1", onde(i, el),
                             "linhas vazias %s e colunas vazias %s usadas como espacador "
                             "— o leitor infere fim de tabela"
                             % (tp["empty_rows"] or "-", tp["empty_cols"] or "-"))
                if not A.get_alt_text(el).strip():
                    rep.fail("G05", "A", "1.1.1", onde(i, el),
                             "tabela sem alt text nem resumo")


# ==========================================================================
# Camada H — hiperlinks
# ==========================================================================
def audit_links(prs, rep: Report):
    for r in ("H01", "H02", "H03"):
        rep.check(r)
    destinos = defaultdict(set)

    for i, slide in enumerate(prs.slides, 1):
        rels = slide.part.rels
        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            for p_el in A.iter_paragraphs(el):
                for r_el in A.iter_runs(p_el):
                    rPr = r_el.find(A.q("a:rPr"))
                    if rPr is None:
                        continue
                    h = rPr.find(A.q("a:hlinkClick"))
                    if h is None:
                        continue
                    rid = h.get(A.q("r:id"))
                    alvo = ""
                    if rid:
                        try:
                            alvo = rels[rid].target_ref
                        except Exception:
                            alvo = rid
                    texto = A.run_text(r_el).strip()
                    if A.RE_URL.match(texto):
                        rep.fail("H01", "A", "2.4.4", onde(i, el),
                                 "texto do link e a URL crua (%r) — o leitor de tela le "
                                 "caractere a caractere" % texto[:50])
                    if texto.lower().strip(" .:;!?") in A.VAGUE_LINK:
                        rep.fail("H02", "A", "2.4.4", onde(i, el),
                                 "texto de link vago: %r" % texto)
                    if texto:
                        destinos[texto.lower()].add(alvo)

    for texto, alvos in destinos.items():
        if len(alvos) > 1:
            rep.fail("H03", "A", "3.2.4", "varios slides",
                     "o texto de link %r aponta para %d destinos diferentes"
                     % (texto[:40], len(alvos)))


# ==========================================================================
# Camada I — midia, movimento e interacao
# ==========================================================================
def audit_media(prs, rep: Report):
    for r in ("I01", "I02", "I06", "I07"):
        rep.check(r)
    # Agregados: uma linha por deck, nao uma por objeto. Com 28 faixas de
    # audio ou 28 janelas de Libras, o relatorio virava inventario.
    janelas_libras, faixas_audio = [], []
    for i, slide in enumerate(prs.slides, 1):
        transicao = slide._element.find(A.q("p:transition"))
        if transicao is not None:
            filhos = [A.local(c) for c in transicao]
            proibidos = [c for c in filhos
                         if c not in ("fade", "cut", "sndAc", "extLst")]
            if proibidos:
                rep.fail("I06", "A", "—", onde(i),
                         "transicao %s — restrinja a fade/cut" % ", ".join(proibidos))

        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True):
            kind = A.media_kind(el)
            if kind == "video":
                # A janela de Libras E a via de acesso; pedir legenda WebVTT
                # nela e exigir tradução da tradução. Num deck com 28 janelas
                # isso enchia o relatorio com 28 linhas sem sentido.
                if (A.shape_name(el) or "").strip().lower() == "janela de libras":
                    janelas_libras.append(i)
                else:
                    rep.unverified("I01", "E", "1.2.2",
                                   "video presente — confirmar legenda WebVTT anexada",
                                   onde(i, el))
            elif kind == "audio":
                faixas_audio.append(i)

            rid, _tip = A.get_hyperlink(el)
            tem_link_run = any(
                r_el.find(A.q("a:rPr")) is not None
                and r_el.find(A.q("a:rPr")).find(A.q("a:hlinkClick")) is not None
                for p_el in A.iter_paragraphs(el) for r_el in A.iter_runs(p_el))
            if rid or tem_link_run:
                box = A.get_xfrm(el)
                if box and (box[2] < EMU_MIN_TARGET or box[3] < EMU_MIN_TARGET):
                    rep.fail("I07", "A", "2.5.8", onde(i, el),
                             "alvo de clique %.2f×%.2f cm, abaixo de 24×24 px CSS "
                             "(0,64×0,64 cm)" % (box[2] / A.EMU_PER_CM, box[3] / A.EMU_PER_CM))


# ==========================================================================
# Camadas J a M — fora do alcance da analise estatica
# ==========================================================================
    _resumo_de_midia(rep, janelas_libras, faixas_audio,
                     len(prs.slides._sldIdLst))
def _resumo_de_midia(rep, janelas_libras, faixas_audio, total):
    if janelas_libras:
        rep.unverified(
            "J01", "E", "LBI · Dec. 5.626/2005",
            "janela de Libras em %d de %d slides — confirmar que a glosa foi "
            "revisada por intérprete (J05)" % (len(janelas_libras), total))
    if faixas_audio:
        rep.unverified(
            "I02", "E", "1.2.1",
            "audiodescrição em %d de %d slides — confirmar a transcrição das "
            "faixas" % (len(faixas_audio), total))


def declare_unverified(rep: Report, tem_janela_libras=False):
    itens = []
    if not tem_janela_libras:
        # com janela embutida, J01 e J02 sao respondidas pelo proprio arquivo:
        # `_resumo_de_midia` relata quantos slides tem janela, e a regra J02 e
        # medida em `audit_design`. Declara-las como inexistentes seria mentir.
        itens += [
            ("J01", "E", "LBI · Dec. 5.626/2005", "existencia de via em Libras para o conteudo"),
            ("J02", "A", "NBR 15290:2016 7.1.3", "dimensoes da janela de Libras (>= 1/2 altura, >= 1/4 largura)"),
        ]
    itens += [
        ("J04", "D", "—", "instrucao de como ligar as Legendas ao Vivo — NAO e propriedade do arquivo: e preferencia da maquina de quem apresenta"),
        ("K02", "A", "NBR 16452:2016", "faixa de audiodescricao e sua transcricao"),
        ("K03", "E", "1.3.2", "navegacao real com leitor de tela (NVDA + Speech Logger)"),
        ("K04", "A", "—", "transcricao linear em .docx acessivel"),
        ("L01", "E", "ISO 14289", "PDF exportado com marcas de estrutura — rodar export_pdfua.py"),
        ("L07", "E", "ISO 14289", "validacao em veraPDF e PAC 2024"),
        ("M01", "—", "—", "verificador nativo do PowerPoint sem Erros nem Avisos"),
        ("M02", "—", "—", "percurso com NVDA (F6, Tab, Ctrl+Shift+S) com log anexado"),
        ("M03", "—", "—", "leitura em escala de cinza do Windows"),
        ("M04", "—", "—", "simulacao de protanopia, deuteranopia e tritanopia"),
        ("M05", "—", "—", "curadoria humana de todo alt text e descricao longa"),
        ("M06", "—", "—", "leitura por pessoa com deficiencia"),
    ]
    for rule, sev, crit, det in itens:
        rep.unverified(rule, sev, crit, det)


# ==========================================================================
# Relatorio
# ==========================================================================
SEV_ORDEM = {"E": 0, "A": 1, "D": 2, "—": 3}


def render_markdown(rep: Report, caminho: str) -> str:
    c = rep.counts()
    erros, avisos = c.get("E", 0), c.get("A", 0)
    dicas, nv = c.get("D", 0), c.get("nao_verificado", 0)
    aprovado = erros == 0 and avisos == 0

    L = []
    L.append("# Relatório de auditoria de acessibilidade\n")
    L.append("**Arquivo:** `%s`  " % os.path.basename(caminho))
    L.append("**Catálogo:** `references/02-catalogo-auditoria.md`  ")
    L.append("**Regras avaliadas automaticamente:** %d  " % len(rep.evaluated))
    L.append(A.carimbo() + "\n")
    L.append("| Severidade | Não conformes |")
    L.append("|---|---|")
    L.append("| Erro | **%d** |" % erros)
    L.append("| Aviso | **%d** |" % avisos)
    L.append("| Dica | %d |" % dicas)
    L.append("| Não verificado | %d |" % nv)
    L.append("")
    if aprovado:
        L.append("> Nenhum Erro ou Aviso automático em aberto. **A entrega ainda não está "
                 "conforme** enquanto a camada M não for cumprida e registrada.\n")
    else:
        L.append("> **Entrega bloqueada:** %d Erros e %d Avisos em aberto.\n" % (erros, avisos))

    naos = [f for f in rep.findings if f["veredito"] == NAO_CONFORME]
    naos.sort(key=lambda f: (SEV_ORDEM.get(f["severidade"], 9), f["regra"]))
    if naos:
        L.append("## Não conformidades\n")
        L.append("| Regra | Sev | Critério | Onde | Detalhe |")
        L.append("|---|---|---|---|---|")
        for f in naos:
            L.append("| %s | %s | %s | %s | %s |" % (
                f["regra"], f["severidade"], f["criterio"],
                f["onde"], f["detalhe"].replace("|", "\\|")))
        L.append("")

    nvs = [f for f in rep.findings if f["veredito"] == NAO_VERIFICADO]
    if nvs:
        L.append("## Não verificado — pendências que exigem outra ferramenta ou um humano\n")
        L.append("Estes itens **não estão aprovados**. Ausência de evidência não é conformidade.\n")
        L.append("| Regra | Sev | Onde | O que falta |")
        L.append("|---|---|---|---|")
        for f in sorted(nvs, key=lambda x: x["regra"]):
            L.append("| %s | %s | %s | %s |" % (
                f["regra"], f["severidade"], f["onde"],
                f["detalhe"].replace("|", "\\|")))
        L.append("")

    L.append("## Regras avaliadas nesta execução\n")
    L.append("`%s`\n" % "`, `".join(sorted(rep.evaluated)))
    return "\n".join(L)


def audit(caminho: str) -> Report:
    prs = Presentation(caminho)
    rep = Report(caminho)
    audit_metadata(prs, rep, caminho)
    audit_run_language(prs, rep)
    audit_structure(prs, rep)
    audit_sections(prs, rep)
    audit_reading_order(prs, rep)
    audit_alt_text(prs, rep)
    audit_contrast_rules(prs, rep)
    audit_typography(prs, rep)
    audit_tables(prs, rep)
    audit_links(prs, rep)
    tem_janela = any(
        (A.shape_name(el) or "").strip().lower() == "janela de libras"
        for slide in prs.slides
        for el in A.iter_shape_elements(slide.shapes._spTree, recurse_groups=True))
    audit_media(prs, rep)
    audit_design.auditar(prs, rep)
    declare_unverified(rep, tem_janela)
    return rep


def main():
    ap = argparse.ArgumentParser(description="Auditor de acessibilidade de .pptx")
    ap.add_argument("pptx")
    ap.add_argument("--json", dest="json_out")
    ap.add_argument("--md", dest="md_out")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    rep = audit(args.pptx)
    md = render_markdown(rep, args.pptx)

    if args.md_out:
        with open(args.md_out, "w", encoding="utf-8") as f:
            f.write(md)
    if args.json_out:
        with open(args.json_out, "w", encoding="utf-8") as f:
            json.dump({"arquivo": args.pptx, "resumo": rep.counts(),
                       "avaliadas": sorted(rep.evaluated),
                       "achados": rep.findings}, f, ensure_ascii=False, indent=2)
    if not args.quiet:
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass
        print(md)

    c = rep.counts()
    return 1 if (c.get("E", 0) or c.get("A", 0)) else 0


if __name__ == "__main__":
    sys.exit(main())
