"""
gen_audiodesc — audiodescricao narrada por slide (regra K02, ABNT NBR 16452).

O roteiro NAO e inventado aqui: ele vem do que ja foi escrito com cuidado no
deck — titulo do slide, texto e a descricao longa das figuras que mora nas
Notas. Gerar um texto novo por cima produziria uma terceira versao do mesmo
conteudo, com risco de divergir das outras duas.

Entrega SEMPRE em par: audio e transcricao. Audiodescricao sem transcricao nao
e auditavel e exclui quem usa linha braille.

    python scripts/gen_audiodesc.py deck.pptx -o audiodescricao/
    python scripts/gen_audiodesc.py deck.pptx -o audiodescricao/ --so-roteiro

Exige ELEVENLABS_API_KEY, exceto com --so-roteiro, que escreve so o texto.
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import a11y_lib as A

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

API = "https://api.elevenlabs.io/v1/text-to-speech/%s"
VOZ_PADRAO = "oi8rgjIfLgJRsQ6rbZh3"   # Amanda Kelly, PT-BR
MODELO = "eleven_multilingual_v2"
PRECO_POR_MIL = 0.30                   # USD, aproximado


def roteiro_do_slide(slide, numero, total) -> str:
    """
    Monta a fala a partir do que ja existe no slide.

    Ordem de anuncio: onde estamos, o que o slide diz, e so entao a descricao
    da figura — do geral ao particular, como pede a NBR 16452.
    """
    titulo_el = next((s for s in A.iter_shape_elements(slide.shapes._spTree)
                      if A.is_title_placeholder(s)), None)
    titulo = ""
    if titulo_el is not None:
        titulo = " ".join(A.paragraph_text(p)
                          for p in A.iter_paragraphs(titulo_el)).strip()

    partes = ["Slide %d de %d. %s." % (numero, total, titulo or "Sem título")]

    corpo = []
    for el in A.iter_shape_elements(slide.shapes._spTree):
        if el is titulo_el or A.is_decorative(el) or A.is_hidden(el):
            continue
        if A.get_tables(el):
            tp = A.table_props(next(iter(A.get_tables(el))))
            alt = A.get_alt_text(el).strip()
            corpo.append("Tabela com %d linhas e %d colunas. %s"
                         % (tp["n_rows"], tp["n_cols"], alt))
            continue
        for p_el in A.iter_paragraphs(el):
            t = A.paragraph_text(p_el).strip()
            if t:
                corpo.append(t if t.endswith((".", "!", "?", ":")) else t + ".")
    if corpo:
        partes.append(" ".join(corpo))

    try:
        if slide.has_notes_slide:
            notas = (slide.notes_slide.notes_text_frame.text or "").strip()
            m = re.search(r"Descrição da figura:\s*(.+)", notas, re.S)
            if m:
                partes.append("Descrição da figura. " + " ".join(m.group(1).split()))
    except Exception:
        pass

    return " ".join(partes)


def sintetizar(texto: str, destino: str, voz: str, chave: str) -> bool:
    import requests

    r = requests.post(
        API % voz,
        headers={"xi-api-key": chave, "Content-Type": "application/json"},
        json={"text": texto, "model_id": MODELO,
              "voice_settings": {"stability": 0.5, "similarity_boost": 0.75,
                                 "speed": 0.95}},
        timeout=180)
    if r.status_code != 200:
        print("      erro %d: %s" % (r.status_code, (r.text or "")[:160]))
        return False
    with open(destino, "wb") as f:
        f.write(r.content)
    return True


def gerar(pptx: str, pasta: str, voz: str = VOZ_PADRAO,
          so_roteiro: bool = False, so_modo: str = "Modo padrão") -> dict:
    prs = Presentation(pptx)
    slides = list(prs.slides)
    os.makedirs(pasta, exist_ok=True)

    # Num deck com 3 modos de cor, o conteudo se repete: narrar so um deles.
    secoes = A.get_sections(prs)
    alvo = range(1, len(slides) + 1)
    if secoes:
        ids = [s.get("id") for s in prs._element.iter(A.q("p:sldId"))]
        for nome, sec_ids in secoes:
            if nome.strip().lower() == so_modo.strip().lower():
                alvo = [ids.index(i) + 1 for i in sec_ids if i in ids]
                break

    chave = (os.environ.get("ELEVENLABS_API_KEY") or "").strip()
    if not so_roteiro and not chave:
        raise RuntimeError(
            "ELEVENLABS_API_KEY ausente. Use --so-roteiro para escrever apenas "
            "o texto, ou defina a variavel de ambiente.")

    linhas_md = ["# Transcrição da audiodescrição\n",
                 "Arquivo: `%s`\n" % os.path.basename(pptx),
                 "Cada bloco corresponde a uma faixa de áudio. A transcrição é "
                 "obrigatória: audiodescrição sem ela não é auditável e exclui "
                 "quem usa linha braille.\n"]
    total_chars = 0
    gerados = 0

    for pos, n in enumerate(alvo, 1):
        slide = slides[n - 1]
        texto = roteiro_do_slide(slide, pos, len(alvo))
        total_chars += len(texto)
        nome = "slide-%02d" % pos
        linhas_md.append("## %s\n" % nome)
        linhas_md.append(texto + "\n")

        if so_roteiro:
            continue
        destino = os.path.join(pasta, nome + ".mp3")
        print("   %s (%d caracteres)" % (nome, len(texto)))
        if sintetizar(texto, destino, voz, chave):
            gerados += 1
            time.sleep(0.4)   # respeita o limite de taxa

    with open(os.path.join(pasta, "transcricao-audiodescricao.md"), "w",
              encoding="utf-8") as f:
        f.write("\n".join(linhas_md))

    return {"faixas": len(alvo), "gerados": gerados, "caracteres": total_chars,
            "custo_usd": round(total_chars / 1000 * PRECO_POR_MIL, 2)}


def main():
    ap = argparse.ArgumentParser(description="Audiodescricao narrada por slide")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--saida", default="audiodescricao")
    ap.add_argument("--voz", default=VOZ_PADRAO)
    ap.add_argument("--so-roteiro", action="store_true",
                    help="escreve so a transcricao, sem chamar a API")
    args = ap.parse_args()

    try:
        r = gerar(args.pptx, args.saida, args.voz, args.so_roteiro)
    except RuntimeError as e:
        print("ERRO: %s" % e)
        return 2

    print("\n%d faixas · %d caracteres · custo estimado US$ %.2f"
          % (r["faixas"], r["caracteres"], r["custo_usd"]))
    if args.so_roteiro:
        print("Somente a transcrição foi escrita (--so-roteiro).")
    else:
        print("%d arquivos .mp3 gerados em %s/" % (r["gerados"], args.saida))
    print("\nA transcrição acompanha o áudio: transcricao-audiodescricao.md")
    print("Entregue os dois. Áudio sem transcrição não é auditável.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
