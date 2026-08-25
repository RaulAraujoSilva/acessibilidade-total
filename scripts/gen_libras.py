"""
gen_libras — camada de Libras (regras J01 a J06).

Tres niveis de entrega, e o script REGISTRA qual foi alcancado. Degradar e
permitido; degradar em silencio, nao.

    Nivel 1  video com janela de Libras, gerado localmente
    Nivel 2  video gerado no portal VLibras Video (passo manual, exige login)
    Nivel 3  glosa + legenda .srt + instrucoes de VLibras Widget/Desktop

    python scripts/gen_libras.py deck.pptx -o libras/

===========================================================================
ESTADO DO NIVEL 1 — o que ja se sabe, para nao reinvestigar
===========================================================================

O renderizador do avatar EXISTE e e invocavel direto, sem RabbitMQ nem
MongoDB. Dentro da imagem `vlibras/translator-video:3.1.0` (436 MB):

    /dist/player/VLibras-Video.x86_64      binario Unity standalone
    /usr/bin/xvfb-run, /usr/bin/Xvfb       display virtual, ja instalados
    /usr/bin/ffmpeg                        montagem dos quadros

A chamada, extraida de /dist/player/playerwrapper.py:

    VLibras-Video.x86_64 --id <tag> --glosapath <arquivo> \\
        --videopath <dir-de-quadros> --width 720 --height 900 \\
        --speed 150 --framerate 24 --avatar icaro --subtitle off \\
        --bundlespath <BUNDLES>

O arquivo de glosa tem o formato `0#GLOSA EM MAIUSCULAS`.

O QUE FALTA: os *bundles* de sinais, apontados por VIDEOMAKER_BUNDLES_DIR.
Eles NAO estao em `translator-video:3.1.0`, nem em `video-core:4.0.0`, nem
aparecem como diretorio em nenhuma das duas — sao servidos em tempo de
execucao pelo servico `dicionario` (imagens `vlibras/dicionario`, ~347 MB).
O proximo passo de quem retomar isto e subir o `dicionario` e descobrir por
onde ele publica os bundles, em vez de comecar do zero.
===========================================================================
"""
from __future__ import annotations

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import a11y_lib as A

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

# Ritmo de fala usado para cronometrar a legenda. Locucao tranquila em
# portugues fica perto disto; audiodescricao costuma ser um pouco mais lenta.
PALAVRAS_POR_MINUTO = 145


def texto_do_slide(slide) -> str:
    partes = []
    for el in A.iter_shape_elements(slide.shapes._spTree):
        if A.is_decorative(el) or A.is_hidden(el):
            continue
        for p_el in A.iter_paragraphs(el):
            t = A.paragraph_text(p_el).strip()
            if t:
                partes.append(t if t.endswith((".", "!", "?", ":")) else t + ".")
    return " ".join(partes)


def _hhmmss(segundos: float) -> str:
    ms = int(round((segundos - int(segundos)) * 1000))
    s = int(segundos)
    return "%02d:%02d:%02d,%03d" % (s // 3600, (s % 3600) // 60, s % 60, ms)


def gerar_srt(blocos, destino: str) -> float:
    """Legenda cronometrada pelo numero de palavras. Devolve a duracao total."""
    linhas, t = [], 0.0
    for i, texto in enumerate(blocos, 1):
        dur = max(2.0, len(texto.split()) / PALAVRAS_POR_MINUTO * 60.0)
        linhas.append(str(i))
        linhas.append("%s --> %s" % (_hhmmss(t), _hhmmss(t + dur)))
        # legenda em duas linhas curtas le melhor que uma longa
        linhas.extend(_quebrar(texto, 42))
        linhas.append("")
        t += dur
    with open(destino, "w", encoding="utf-8") as f:
        f.write("\n".join(linhas))
    return t


def _quebrar(texto, largura):
    saida, atual = [], ""
    for p in texto.split():
        if len(atual) + len(p) + 1 > largura:
            saida.append(atual)
            atual = p
        else:
            atual = (atual + " " + p).strip()
        if len(saida) == 2:
            break
    if atual and len(saida) < 2:
        saida.append(atual)
    return saida or [texto[:largura]]


def gerar_glosa(blocos):
    """
    Traducao para glosa com vlibras-translate, se disponivel.

    Glosa automatica NAO e traducao revisada: erra concordancia espacial e
    classificadores. Serve de insumo para intérprete, nunca de entrega final
    (regra J05).
    """
    try:
        from vlibras_translate import translation
    except ImportError:
        return None, ("vlibras-translate ausente. Instale com: "
                      "pip install vlibras-translate")
    try:
        tradutor = translation.Translation()
        return [tradutor.rooting(b) for b in blocos], None
    except Exception as e:
        return None, "vlibras-translate falhou: %s" % e


def gerar(pptx: str, pasta: str, so_modo: str = "Modo padrão") -> dict:
    prs = Presentation(pptx)
    slides = list(prs.slides)
    os.makedirs(pasta, exist_ok=True)

    secoes = A.get_sections(prs)
    alvo = list(range(1, len(slides) + 1))
    if secoes:
        ids = [s.get("id") for s in prs._element.iter(A.q("p:sldId"))]
        for nome, sec_ids in secoes:
            if nome.strip().lower() == so_modo.strip().lower():
                alvo = [ids.index(i) + 1 for i in sec_ids if i in ids]
                break

    blocos = [texto_do_slide(slides[n - 1]) for n in alvo]
    blocos = [b for b in blocos if b.strip()]

    srt = os.path.join(pasta, "roteiro-libras.srt")
    duracao = gerar_srt(blocos, srt)

    glosa, erro_glosa = gerar_glosa(blocos)
    if glosa:
        with open(os.path.join(pasta, "glosa.txt"), "w", encoding="utf-8") as f:
            for i, g in enumerate(glosa, 1):
                f.write("# bloco %d\n%s\n\n" % (i, g))

    nivel = 3
    with open(os.path.join(pasta, "LEIA-ME.md"), "w", encoding="utf-8") as f:
        f.write(_leia_me(os.path.basename(pptx), len(blocos), duracao,
                         bool(glosa), erro_glosa))

    return {"blocos": len(blocos), "duracao_s": duracao, "nivel": nivel,
            "glosa": bool(glosa), "erro_glosa": erro_glosa}


def _leia_me(arquivo, n, duracao, tem_glosa, erro_glosa):
    return """# Camada de Libras — o que foi entregue

**Arquivo de origem:** `%s`
**Blocos de conteúdo:** %d · **duração estimada da narração:** %d min %02d s

## Nível alcançado: 3 de 3 — roteiro e legenda, sem vídeo de avatar

Isto está registrado de propósito. A regra J01 continua **não atendida**: não há
janela de Libras neste material. O que existe é o insumo para produzi-la.

| Arquivo | O que é |
|---|---|
| `roteiro-libras.srt` | Legenda cronometrada, pronta para o portal VLibras Vídeo |
| `glosa.txt` | Tradução automática para glosa%s |

## Como chegar ao nível 2 (vídeo, hoje)

1. Abra https://video.vlibras.gov.br/ e entre com a conta gov.br.
2. Envie um vídeo da apresentação e o `roteiro-libras.srt`.
3. O portal devolve o vídeo com a janela de Libras.

Limites do portal: `.mp4` até 500 MB e `.srt` até 600 palavras.

## Como chegar ao nível 1 (local, sem login)

O renderizador do avatar é um binário Unity que **já está** na imagem
`vlibras/translator-video:3.1.0`, junto com Xvfb e ffmpeg, e aceita ser chamado
direto — sem RabbitMQ nem MongoDB. A chamada está documentada no cabeçalho de
`scripts/gen_libras.py`.

O que falta são os *bundles* de sinais (`VIDEOMAKER_BUNDLES_DIR`), que não estão
naquela imagem nem em `vlibras/video-core:4.0.0`. Eles são servidos pelo serviço
`dicionario`. Quem retomar deve começar por aí, e não do zero.

## Antes de publicar

Glosa automática erra concordância espacial e classificadores. Ela é insumo para
intérprete, **não** entrega final (regra J05). E os parâmetros da janela — altura
de pelo menos metade da tela, largura de pelo menos um quarto — estão na
ABNT NBR 15290:2016, item 7.1.3, com a ressalva de que aplicá-los a um slide é
analogia, porque a norma regula televisão.
""" % (arquivo, n, int(duracao // 60), int(duracao % 60),
       "" if tem_glosa else " — **nao gerada**: " + (erro_glosa or "motivo desconhecido"))


def main():
    ap = argparse.ArgumentParser(description="Camada de Libras do material")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--saida", default="libras")
    args = ap.parse_args()

    r = gerar(args.pptx, args.saida)
    print("blocos: %d · narração estimada: %d min %02d s"
          % (r["blocos"], r["duracao_s"] // 60, r["duracao_s"] % 60))
    print("legenda: roteiro-libras.srt")
    if r["glosa"]:
        print("glosa:   glosa.txt (automática, precisa de revisão — regra J05)")
    else:
        print("glosa:   NÃO gerada — %s" % r["erro_glosa"])
    print("\nNível alcançado: 3 de 3 (roteiro e legenda, sem vídeo de avatar).")
    print("A regra J01 segue NÃO ATENDIDA e assim consta no relatório.")
    print("Veja %s/LEIA-ME.md para o caminho até os níveis 2 e 1." % args.saida)
    return 0


if __name__ == "__main__":
    sys.exit(main())
