"""
gerar_modelo — produz assets/modelo-acessivel.pptx.

POR QUE ISTO EXISTE: enquanto o construtor herdava o template padrao do Office,
quem mandava no layout era a Microsoft, nao o projeto. O resultado media-se:

  - o layout "Section Header" traz cap="all" no slideLayout3.xml, e os titulos
    de secao saiam em CAIXA ALTA — o deck violava a regra F06 que ele ensina;
  - nesse mesmo layout o corpo fica em y=8,07 cm e o titulo em y=12,24 cm, ou
    seja, o apoio aparece ACIMA do titulo;
  - titulo sem posicao explicita = posicao decidida pelo template.

Aqui os 11 layouts existentes sao REAPROVEITADOS e reposicionados na grade do
projeto (scripts/grade.py). Reaproveitar em vez de criar partes novas evita
mexer nas relacoes OPC, que e onde um .pptx costuma quebrar.

    python scripts/gerar_modelo.py
"""
from __future__ import annotations

import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

import a11y_lib as A
import grade as G

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "assets", "modelo-acessivel.pptx")
PALETA = os.path.join(RAIZ, "assets", "paleta-okabe-ito.json")

FONTE = "Calibri"

# --------------------------------------------------------------------------
# Os layouts do projeto: indice do template -> nome e geometria
#
# A geometria e um mapa {idx_do_placeholder: (coluna, colunas, y, altura)}.
# Placeholders de data, rodape e numero de slide ficam como estao: o
# python-pptx nao os copia para o slide, entao nao atrapalham.
# --------------------------------------------------------------------------
T, TH = G.TITULO_Y, G.TITULO_H
C, CH = G.CONTEUDO_Y, G.CONTEUDO_H
meia = CH // 2

LAYOUTS = {
    0: ("Capa", {
        # O subtitulo da capa leva 4 linhas (chamada, disciplina, programa,
        # autor) e a 4a transbordava para dentro da faixa do rodape: a CAIXA
        # terminava aos 15,6 cm, mas o TEXTO nao. Titulo mais alto e caixa de
        # subtitulo com folga real.
        0: (0, 8, G.cm(4.2), G.cm(3.4)),          # titulo grande
        1: (0, 8, G.cm(8.2), G.cm(8.6)),          # subtitulo
    }),
    1: ("Título e conteúdo", {
        0: (0, 12, T, TH),
        1: (0, 12, C, CH),
    }),
    2: ("Seção", {
        0: (0, 10, G.cm(6.6), G.cm(3.4)),         # titulo ACIMA do apoio
        1: (0, 10, G.cm(10.2), G.cm(2.4)),        # apoio logo abaixo
    }),
    3: ("Duas colunas", {
        0: (0, 12, T, TH),
        1: (0, 6, C, CH),
        2: (6, 6, C, CH),
    }),
    4: ("Comparação", {
        0: (0, 12, T, TH),
        1: (0, 6, C, G.cm(1.5)),
        2: (0, 6, C + G.cm(1.7), CH - G.cm(1.7)),
        3: (6, 6, C, G.cm(1.5)),
        4: (6, 6, C + G.cm(1.7), CH - G.cm(1.7)),
    }),
    5: ("Título e área livre", {
        0: (0, 12, T, TH),
    }),
    6: ("Em branco", {}),                          # nome mantido: regra B06
    7: ("Destaque", {
        0: (0, 12, T, TH),
        1: (0, 7, C, CH),
        2: (7, 5, C, CH),
    }),
    8: ("Figura com legenda", {
        0: (0, 12, T, TH),
        1: (0, 7, C, CH),
        2: (7, 5, C, CH),
    }),
    9: ("Cartões", {
        0: (0, 12, T, TH),
        1: (0, 12, C, CH),
    }),
    10: ("Citação", {
        0: (0, 10, G.cm(5.6), G.cm(6.4)),
        1: (0, 10, G.cm(12.6), G.cm(2.4)),
    }),
}


# --------------------------------------------------------------------------
# Mobilia decorativa: filete de acento e faixa do rodape
#
# Ela MORA NO LAYOUT, nao no slide. Desenhada slide a slide, a forma existe na
# arvore do slide e o Verificador de Acessibilidade nativo a lista no painel de
# ordem de leitura — mesmo marcada como decorativa. Foi o que ele acusou em
# 26/08/2026: "faixa decorativa" e "faixa de acento decorativa" aparecendo em
# todos os slides.
#
# No layout, ela e cromo herdado: nao entra no spTree do slide, nao aparece no
# painel, e nao ha o que reordenar. A cor vem do TEMA (accent1), entao cada
# modo de cor reescreve o clrScheme e a mobilia acompanha sem ser redesenhada.
# --------------------------------------------------------------------------
FILETE_FINO = G.cm(0.16)
FILETE_GROSSO = G.cm(0.3)

MOBILIA = {
    "Capa":               [("filete", 0, 4, G.cm(3.2), FILETE_GROSSO)],
    "Seção":              [("filete", 0, 3, G.cm(5.6), FILETE_GROSSO)],
    "Citação":            [("filete", 0, 2, G.cm(4.6), FILETE_GROSSO)],
    "Em branco":          [],
}
_FILETE_PADRAO = ("filete", 0, 2, T + TH + G.cm(0.2), FILETE_FINO)


def mobiliar(layout, nome) -> int:
    """
    Poe o filete e a faixa do rodape no LAYOUT, com cor de tema.

    `LayoutShapes` nao expoe `add_shape` — o python-pptx so monta formas em
    slide. Como sao dois retangulos simples, o XML vai a mao.
    """
    itens = MOBILIA.get(nome, [_FILETE_PADRAO])
    spTree = layout._element.find(A.q("p:cSld")).find(A.q("p:spTree"))
    usados = [int(c.get("id")) for c in spTree.iter(A.q("p:cNvPr"))
              if (c.get("id") or "").isdigit()]
    proximo = (max(usados) + 1) if usados else 2
    postos = 0

    def _forma(x, y, larg, alt, rotulo, ident):
        xml = (
            '<p:sp xmlns:p="%s" xmlns:a="%s">'
            '<p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr/><p:nvPr/>'
            '</p:nvSpPr><p:spPr>'
            '<a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            '<a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
            '<a:ln><a:noFill/></a:ln>'
            '</p:spPr></p:sp>'
            % (A.NS["p"], A.NS["a"], ident, rotulo, x, y, larg, alt))
        el = etree.fromstring(xml)
        # decorativo tambem aqui: no layout ele ja e cromo herdado, mas a marca
        # deixa a intencao explicita para quem abrir o modelo
        A.set_decorative(el, True)
        spTree.append(el)
        return el

    for item in itens:
        if item[0] == "filete":
            _, coluna, n, y, esp = item
            _forma(G.x(coluna), y, G.larg(n), esp, "Filete de acento", proximo)
            proximo += 1
            postos += 1

    if nome != "Em branco":
        _forma(0, G.ALTURA - G.FAIXA_H, G.LARGURA, G.FAIXA_H,
               "Faixa estética do rodapé", proximo)
        proximo += 1
        postos += 1
    return postos


def limpar_caixa_alta(el) -> int:
    """
    Remove cap="all" e cap="small" de qualquer defRPr/rPr.

    E a origem do defeito D4: o titulo saia em CAIXA ALTA por heranca, o texto
    no XML continuava em caixa mista, e por isso a regra F06 nao via nada.
    """
    n = 0
    for tag in ("a:defRPr", "a:rPr", "a:endParaRPr"):
        for r in el.iter(A.q(tag)):
            if r.get("cap") in ("all", "small"):
                del r.attrib["cap"]
                n += 1
    return n


def horizontalizar(el) -> int:
    """Tira o texto vertical dos layouts reaproveitados para Cartões e Citação."""
    n = 0
    for bp in el.iter(A.q("a:bodyPr")):
        if bp.get("vert") and bp.get("vert") != "horz":
            del bp.attrib["vert"]
            n += 1
    return n


def tipografia_do_master(master):
    """Escala, alinhamento e entrelinha no p:txStyles — a base de tudo."""
    styles = master._element.find(A.q("p:txStyles"))
    if styles is None:
        return
    plano = {
        "p:titleStyle": [(G.TIPO["titulo"], G.ENTRELINHA_TITULO)],
        "p:bodyStyle": [(G.TIPO["corpo"], G.ENTRELINHA_CORPO),
                        (G.TIPO["corpo_denso"], G.ENTRELINHA_CORPO),
                        (G.TIPO["apoio"], G.ENTRELINHA_CORPO),
                        (G.TIPO["apoio"], G.ENTRELINHA_CORPO),
                        (G.TIPO["apoio"], G.ENTRELINHA_CORPO)],
        "p:otherStyle": [(G.TIPO["apoio"], G.ENTRELINHA_CORPO)],
    }
    for nome, niveis in plano.items():
        st = styles.find(A.q(nome))
        if st is None:
            continue
        for i, (pt, lnspc) in enumerate(niveis, 1):
            lvl = st.find(A.q("a:lvl%dpPr" % i))
            if lvl is None:
                continue
            lvl.set("algn", "l")                       # nunca justificado (F04)
            for antigo in lvl.findall(A.q("a:lnSpc")):
                lvl.remove(antigo)
            ln = etree.Element(A.q("a:lnSpc"))
            etree.SubElement(ln, A.q("a:spcPct")).set("val", str(lnspc))
            lvl.insert(0, ln)
            d = lvl.find(A.q("a:defRPr"))
            if d is None:
                d = etree.SubElement(lvl, A.q("a:defRPr"))
            d.set("sz", str(pt * 100))
            latin = d.find(A.q("a:latin"))
            if latin is None:
                latin = etree.SubElement(d, A.q("a:latin"))
            latin.set("typeface", FONTE)


def paleta_no_tema(master):
    """
    Grava a paleta do projeto no a:clrScheme.

    Assim, o que herdar cor do tema ja cai na paleta cega-segura, e o auditor
    de contraste resolve para a cor certa em vez de topar com o azul do Office.
    """
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    with open(PALETA, encoding="utf-8") as f:
        p = json.load(f)
    marcas = {s["nome"]: s["texto_min_4_5_1"] for s in p["series"]}

    cores = {
        "dk1": p["textos"]["sobre_claro"], "lt1": p["fundos"]["padrao"],
        "dk2": "#333333", "lt2": "#FFFFFF",
        "accent1": marcas["azul"], "accent2": marcas["laranja"],
        "accent3": marcas["verde_azulado"], "accent4": marcas["roxo_avermelhado"],
        "accent5": marcas["vermelhao"], "accent6": marcas["cinza"],
        "hlink": marcas["azul"], "folHlink": marcas["roxo_avermelhado"],
    }
    try:
        parte = master.part.part_related_by(RT.THEME)
        raiz = etree.fromstring(parte.blob)
    except Exception:
        return 0
    scheme = raiz.find(".//" + A.q("a:clrScheme"))
    if scheme is None:
        return 0
    n = 0
    for filho in scheme:
        nome = A.local(filho)
        if nome not in cores:
            continue
        for antigo in list(filho):
            filho.remove(antigo)
        etree.SubElement(filho, A.q("a:srgbClr")).set(
            "val", cores[nome].lstrip("#").upper())
        n += 1
    parte._blob = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8",
                                 standalone=True)
    # python-pptx guarda partes XML como blob; forcar a reescrita
    try:
        parte.blob  # noqa
        parte._blob_override = None
    except Exception:
        pass
    return n


def montar_cartoes(layout, quantos=4):
    """
    Clona o placeholder de corpo para ter um por cartao.

    Sem isto, cartao viraria caixa de texto solta — e caixa solta reprova em
    B05, com razao: ela nao carrega semantica. Placeholder de verdade mantem a
    estrutura e ainda entra na ordem de leitura corretamente.
    """
    import copy

    corpo = None
    for ph in layout.placeholders:
        if ph.placeholder_format.idx == 1:
            corpo = ph._element
            break
    if corpo is None:
        return 0

    spTree = corpo.getparent()
    maior_id = max(int(c.get("id"))
                   for c in spTree.iter(A.q("p:cNvPr")) if c.get("id"))
    criados = 0
    for k in range(1, quantos):
        novo_el = copy.deepcopy(corpo)
        cNvPr = A.get_cNvPr(novo_el)
        maior_id += 1
        cNvPr.set("id", str(maior_id))
        cNvPr.set("name", "Cartão %d" % (k + 1))
        ph = A.get_ph(novo_el)
        ph.set("idx", str(1 + k))
        spTree.append(novo_el)
        criados += 1

    # dispor os cartoes numa linha, 3 colunas cada
    for k in range(quantos):
        for p_ in layout.placeholders:
            if p_.placeholder_format.idx == 1 + k:
                p_.left = G.x(k * 3)
                p_.top = G.CONTEUDO_Y
                p_.width = G.larg(3)
                p_.height = G.cm(7.4)
    return criados


def aplicar_geometria(layout, mapa):
    ajustados = 0
    for ph in layout.placeholders:
        idx = ph.placeholder_format.idx
        if idx not in mapa:
            continue
        coluna, n, y, h = mapa[idx]
        ph.left, ph.top = G.x(coluna), y
        ph.width, ph.height = G.larg(n), h
        ajustados += 1
    return ajustados


def gerar(saida: str = SAIDA) -> dict:
    prs = Presentation()
    prs.slide_width, prs.slide_height = G.LARGURA, G.ALTURA
    master = prs.slide_masters[0]

    caps = vert = geo = mob = 0
    for i, (nome, mapa) in LAYOUTS.items():
        layout = prs.slide_layouts[i]
        el = layout._element
        caps += limpar_caixa_alta(el)
        vert += horizontalizar(el)
        geo += aplicar_geometria(layout, mapa)
        mob += mobiliar(layout, nome)
        if nome == "Cartões":
            montar_cartoes(layout)
        cSld = el.find(A.q("p:cSld"))
        if cSld is not None:
            cSld.set("name", nome)

    caps += limpar_caixa_alta(master._element)
    tipografia_do_master(master)
    cores = paleta_no_tema(master)

    os.makedirs(os.path.dirname(saida), exist_ok=True)
    prs.save(saida)
    return {"caps_removidos": caps, "vert_corrigidos": vert,
            "mobilia_nos_layouts": mob,
            "placeholders_posicionados": geo, "cores_no_tema": cores}


def validar(caminho: str) -> list:
    """Abre o modelo e confere que ele nao guarda mais os defeitos conhecidos."""
    problemas = []
    prs = Presentation(caminho)

    if (prs.slide_width, prs.slide_height) != (G.LARGURA, G.ALTURA):
        problemas.append("tamanho de slide fora de 16:9")

    for i, (nome, mapa) in LAYOUTS.items():
        layout = prs.slide_layouts[i]
        if layout.name != nome:
            problemas.append("layout %d chama-se %r, esperado %r"
                             % (i, layout.name, nome))
        caps = [r.get("cap") for r in layout._element.iter(A.q("a:defRPr"))
                if r.get("cap") in ("all", "small")]
        if caps:
            problemas.append("layout %r ainda tem cap=%s" % (nome, caps))
        # titulo tem de estar ACIMA do corpo
        pos = {}
        for ph in layout.placeholders:
            pos[ph.placeholder_format.idx] = ph.top or 0
        if nome != "Cartões" and 0 in mapa and 1 in mapa and 0 in pos and 1 in pos:
            if pos[0] > pos[1]:
                problemas.append("layout %r: corpo acima do titulo" % nome)
        for idx, (_c, _n, y, h) in mapa.items():
            if y + h > G.RODAPE_Y + G.TOL_GRADE:
                problemas.append("layout %r: placeholder %d invade o rodape"
                                 % (nome, idx))

    # o modelo tem de conseguir gerar um slide de cada layout
    for i in LAYOUTS:
        try:
            prs.slides.add_slide(prs.slide_layouts[i])
        except Exception as e:
            problemas.append("layout %d nao gera slide: %s" % (i, e))
    return problemas


def main():
    r = gerar()
    print("modelo gerado: %s" % SAIDA)
    print("  %d atributos de CAIXA ALTA removidos" % r["caps_removidos"])
    print("  %d textos verticais horizontalizados" % r["vert_corrigidos"])
    print("  %d placeholders posicionados na grade" % r["placeholders_posicionados"])
    print("  %d cores gravadas no tema" % r["cores_no_tema"])

    print("\nvalidando...")
    problemas = validar(SAIDA)
    if problemas:
        print("  MODELO INVALIDO:")
        for p in problemas:
            print("    - %s" % p)
        return 1
    print("  ok: nenhuma caixa alta, titulo sempre acima do corpo, nada no rodapé,")
    print("      e todos os layouts geram slide.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
