"""
gen_transcricao — transcricao linear do deck em .docx acessivel (regra K04).

Por que isto nao e um extra: as DESCRICOES LONGAS moram nas Notas do orador, e
as notas EVAPORAM na exportacao padrao para PDF. Se o PDF for o unico
entregavel, a descricao longa de cada figura se perde. A transcricao e o que
preserva essa camada fora do .pptx.

Serve tambem a quem prefere ler o conteudo de forma linear em vez de navegar
slide a slide, e a quem usa linha braille.

    python scripts/gen_transcricao.py deck.pptx -o transcricao.docx
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import a11y_lib as A

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def _tabela_para_docx(doc, tbl, legenda):
    linhas = tbl.findall(A.q("a:tr"))
    if not linhas:
        return
    ncols = max(len(tr.findall(A.q("a:tc"))) for tr in linhas)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    tp = A.table_props(tbl)
    for ri, tr in enumerate(linhas):
        celulas = tr.findall(A.q("a:tc"))
        linha = t.add_row()
        for ci in range(ncols):
            txt = A._tc_text(celulas[ci]).strip() if ci < len(celulas) else ""
            cel = linha.cells[ci]
            cel.text = txt
            if ri == 0 and tp["first_row"]:
                for p in cel.paragraphs:
                    for r in p.runs:
                        r.bold = True
    if legenda:
        p = doc.add_paragraph(legenda)
        p.style = "Caption"


def gerar(pptx: str, saida: str) -> dict:
    from docx import Document
    from docx.shared import Pt

    prs = Presentation(pptx)
    cp = prs.core_properties
    doc = Document()

    # tipografia acessivel tambem no .docx
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5

    doc.core_properties.title = cp.title or os.path.basename(pptx)
    doc.core_properties.author = cp.author or ""
    doc.core_properties.language = cp.language or "pt-BR"
    doc.core_properties.subject = cp.subject or ""

    doc.add_heading(cp.title or "Transcrição da apresentação", level=0)
    doc.add_paragraph(
        "Transcrição linear de %s. Cada seção corresponde a um slide, na ordem "
        "em que aparecem. As descrições longas das figuras estão reproduzidas "
        "aqui porque as Anotações do orador não sobrevivem à exportação para "
        "PDF." % os.path.basename(pptx))

    contagem = {"slides": 0, "figuras": 0, "tabelas": 0, "descricoes_longas": 0}

    for i, slide in enumerate(prs.slides, 1):
        titulo_el = next((s for s in A.iter_shape_elements(slide.shapes._spTree)
                          if A.is_title_placeholder(s)), None)
        titulo = ""
        if titulo_el is not None:
            titulo = " ".join(A.paragraph_text(p)
                              for p in A.iter_paragraphs(titulo_el)).strip()
        doc.add_heading(titulo or "Slide %d (sem título)" % i, level=1)
        contagem["slides"] += 1

        for el in A.iter_shape_elements(slide.shapes._spTree):
            if el is titulo_el or A.is_decorative(el) or A.is_hidden(el):
                continue

            tabelas = A.get_tables(el)
            if tabelas:
                alt = A.get_alt_text(el).strip()
                for tbl in tabelas:
                    _tabela_para_docx(doc, tbl, alt)
                    contagem["tabelas"] += 1
                continue

            textos = [A.paragraph_text(p).strip()
                      for p in A.iter_paragraphs(el)]
            textos = [t for t in textos if t]
            if textos:
                for t in textos:
                    doc.add_paragraph(t, style="List Bullet"
                                      if len(textos) > 1 else None)
                continue

            alt = A.get_alt_text(el).strip()
            if alt:
                doc.add_heading("Figura", level=2)
                doc.add_paragraph(alt)
                contagem["figuras"] += 1

        try:
            notas = ""
            if slide.has_notes_slide:
                notas = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notas:
                doc.add_heading("Descrição detalhada e notas", level=2)
                for linha in notas.splitlines():
                    if linha.strip():
                        doc.add_paragraph(linha.strip())
                contagem["descricoes_longas"] += 1
        except Exception:
            pass

    doc.save(saida)
    return contagem


def main():
    ap = argparse.ArgumentParser(
        description="Gera a transcricao linear do deck em .docx acessivel")
    ap.add_argument("pptx")
    ap.add_argument("-o", "--saida")
    args = ap.parse_args()

    saida = args.saida or os.path.splitext(args.pptx)[0] + "-transcricao.docx"
    try:
        c = gerar(args.pptx, saida)
    except ImportError:
        print("python-docx ausente. Rode: pip install python-docx")
        return 2

    print("gerado: %s" % saida)
    print("  %d slides · %d figuras descritas · %d tabelas · %d blocos de "
          "descrição longa" % (c["slides"], c["figuras"], c["tabelas"],
                               c["descricoes_longas"]))
    if c["figuras"] and not c["descricoes_longas"]:
        print("\n  ATENÇÃO: há figuras mas nenhuma descrição longa nas Notas "
              "(regra D07).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
