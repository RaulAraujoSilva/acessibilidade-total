"""
simular_leitura — escreve o que um leitor de tela anunciaria, na ordem em que
anunciaria.

Serve para uma pessoa vidente ENXERGAR a ordem de leitura sem instalar nada.
Le a mesma fonte que o leitor de tela le: a ordem do p:spTree, os textos
alternativos, a marca de decorativo, os cabecalhos de tabela e as notas.

    python scripts/simular_leitura.py deck.pptx
    python scripts/simular_leitura.py deck.pptx --slide 3 --md leitura.md

LIMITE HONESTO: isto e um MODELO do comportamento, nao o comportamento. Nao
substitui a regra K03, que exige o percurso real com NVDA ou JAWS. O que ele
faz e reduzir a chance de surpresa e tornar a ordem discutivel antes do teste.
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation

import a11y_lib as A

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def _texto_da_forma(el) -> list[str]:
    linhas = []
    for p_el in A.iter_paragraphs(el):
        t = A.paragraph_text(p_el).strip()
        if t:
            pp = A.paragraph_props(p_el)
            marcador = "marcador, " if pp["bullet"] in ("buChar", "buAutoNum") else ""
            linhas.append(marcador + t)
    return linhas


def _links_da_forma(el, slide) -> list[str]:
    out = []
    for p_el in A.iter_paragraphs(el):
        for r_el in A.iter_runs(p_el):
            rPr = r_el.find(A.q("a:rPr"))
            if rPr is None:
                continue
            if rPr.find(A.q("a:hlinkClick")) is not None:
                out.append(A.run_text(r_el).strip())
    return out


def _anuncia_tabela(el, tbl) -> list[str]:
    tp = A.table_props(tbl)
    fala = ["tabela, %d linhas, %d colunas" % (tp["n_rows"], tp["n_cols"])]
    linhas = tbl.findall(A.q("a:tr"))
    cabecalhos = []
    if tp["first_row"] and linhas:
        cabecalhos = [A._tc_text(tc).strip() for tc in linhas[0].findall(A.q("a:tc"))]
        fala.append("  linha de cabecalho: " + " | ".join(cabecalhos))
    else:
        fala.append("  SEM linha de cabecalho — o leitor nao consegue "
                    "recontextualizar as celulas (regra G01)")
    corpo = linhas[1:] if tp["first_row"] else linhas
    for ri, tr in enumerate(corpo, 1):
        celulas = tr.findall(A.q("a:tc"))
        partes = []
        for ci, tc in enumerate(celulas):
            v = A._tc_text(tc).strip()
            if cabecalhos and ci < len(cabecalhos) and cabecalhos[ci]:
                partes.append("%s: %s" % (cabecalhos[ci], v or "vazio"))
            else:
                partes.append(v or "vazio")
        fala.append("  linha %d — %s" % (ri, ", ".join(partes)))
    return fala


TIPO_ANUNCIADO = {
    "pic": "imagem",
    "graphicFrame": "objeto",
    "grpSp": "grupo",
    "cxnSp": "conector",
}


def simular_slide(slide, numero, total, mostrar_ignorados=False) -> list[str]:
    fala = []
    titulo_el = next((s for s in A.iter_shape_elements(slide.shapes._spTree)
                      if A.is_title_placeholder(s)), None)
    titulo = ""
    if titulo_el is not None:
        titulo = " ".join(_texto_da_forma(titulo_el)).strip()

    if titulo:
        fala.append('Slide %d de %d: "%s"' % (numero, total, titulo))
    else:
        fala.append("Slide %d de %d: SEM TITULO — no indice de titulos este slide "
                    "e indistinguivel dos outros (regra B01)" % (numero, total))

    for el in A.iter_shape_elements(slide.shapes._spTree):
        nome = A.shape_name(el) or A.local(el)

        if A.is_decorative(el):
            if mostrar_ignorados:
                fala.append("   (silencio — %s marcado como decorativo)" % nome)
            continue
        if A.is_hidden(el):
            if mostrar_ignorados:
                fala.append("   (oculto no Painel de Selecao — %s)" % nome)
            continue
        if el is titulo_el:
            continue

        tabelas = A.get_tables(el)
        if tabelas:
            alt = A.get_alt_text(el).strip()
            if alt:
                fala.append("   %s" % alt)
            for tbl in tabelas:
                fala.extend("   " + l for l in _anuncia_tabela(el, tbl))
            continue

        textos = _texto_da_forma(el)
        links = _links_da_forma(el, slide)
        if textos:
            for t in textos:
                if t in links or any(t.endswith(l) for l in links if l):
                    fala.append("   link, %s" % t)
                else:
                    fala.append("   %s" % t)
            continue

        alt = A.get_alt_text(el).strip()
        tipo = TIPO_ANUNCIADO.get(A.local(el), "objeto")
        if alt:
            fala.append("   %s, %s" % (tipo, alt))
        else:
            fala.append("   %s, SEM TEXTO ALTERNATIVO — o leitor anuncia so "
                        '"%s" e a informacao se perde (regra D01)' % (tipo, tipo))

    try:
        if slide.has_notes_slide:
            notas = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notas:
                fala.append("   [Ctrl+Shift+S — notas do orador]")
                for linha in notas.splitlines():
                    if linha.strip():
                        fala.append("      %s" % linha.strip())
    except Exception:
        pass

    return fala


def main():
    ap = argparse.ArgumentParser(
        description="Simula o que um leitor de tela anunciaria")
    ap.add_argument("pptx")
    ap.add_argument("--slide", type=int, help="simula apenas este slide")
    ap.add_argument("--md", dest="md_out", help="grava a simulacao em Markdown")
    ap.add_argument("--ignorados", action="store_true",
                    help="mostra tambem o que fica em silencio")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    slides = list(prs.slides)
    total = len(slides)

    blocos = []
    for i, slide in enumerate(slides, 1):
        if args.slide and i != args.slide:
            continue
        blocos.append(simular_slide(slide, i, total, args.ignorados))

    cabecalho = [
        "# Leitura simulada — %s" % os.path.basename(args.pptx),
        "",
        "O que um leitor de tela anunciaria, na ordem em que anunciaria.",
        "",
        "> Isto e um **modelo** do comportamento, nao o comportamento. Nao "
        "substitui a regra K03, que exige o percurso real com NVDA. Serve para "
        "discutir a ordem de leitura antes do teste — e para quem enxerga "
        "entender o que quem nao enxerga vai receber.",
        "",
    ]
    corpo = []
    for bloco in blocos:
        corpo.append("```")
        corpo.extend(bloco)
        corpo.append("```")
        corpo.append("")

    md = "\n".join(cabecalho + corpo)
    if args.md_out:
        with open(args.md_out, "w", encoding="utf-8") as f:
            f.write(md)
        print("gravado: %s" % args.md_out)
    else:
        for bloco in blocos:
            print()
            for linha in bloco:
                print(linha)
        print()
    return 0


if __name__ == "__main__":
    sys.exit(main())
