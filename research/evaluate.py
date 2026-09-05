"""Corpus sintético público. Compara trabalho realizado, sem inferir compreensão humana."""
import sys,os,json,time,copy,re,hashlib,statistics
from pathlib import Path
root=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root/'web'))
from studio.content import extract,compare,fingerprint
from studio.exporters import export_all
from studio.providers import adapt_block,models,estimate
from docx import Document
from docx.shared import Cm
from pptx import Presentation
from pptx.util import Inches
from PIL import Image,ImageDraw
import fitz
out=root/'research/evaluation-output';out.mkdir(parents=True,exist_ok=True)
corpus=root/'research/corpus';corpus.mkdir(parents=True,exist_ok=True)
key=os.environ.get('OPENROUTER_API_KEY','')
model_id='openai/gpt-4.1-mini'
catalog=models();model=next(m for m in catalog if m['id']==model_id)
image=Image.new('RGB',(600,220),'white');draw=ImageDraw.Draw(image)
draw.rectangle((100,40,340,90),fill='#164b78');draw.text((10,60),'Grupo A',fill='black');draw.text((350,60),'24',fill='black')
draw.rectangle((100,130,460,180),fill='#9e5d08');draw.text((10,150),'Grupo B',fill='black');draw.text((470,150),'36',fill='black')
img=corpus/'figura.png';image.save(img)
scenarios=[('Atendimento','O prazo para resposta é de 30 dias. O recurso suspende esse prazo, mas não garante a aprovação.','Foram recebidas 60 solicitações: 24 do grupo A e 36 do grupo B. A contagem é descritiva e não demonstra a causa da diferença.'),('Capacitação','A oficina oferece 20 vagas e exige inscrição até 15 de outubro. A participação é gratuita, exceto o transporte.','O resultado preliminar indica 80% de conclusão entre 10 participantes. Esse resultado não pode ser generalizado para toda a turma.')]
for i,(title,a,b) in enumerate(scenarios,1):
    if all((corpus/f'{prefix}{i}.{ext}').exists() for prefix,ext in [('D','docx'),('P','pptx'),('F','pdf')]):continue
    doc=Document();doc.add_heading(title,1);doc.add_paragraph(a);doc.add_paragraph(b)
    t=doc.add_table(rows=3,cols=2)
    rows=[['Grupo','Solicitações'],['A','24'],['B','36']]
    for ri,row in enumerate(rows):
        for ci,value in enumerate(row):t.cell(ri,ci).text=value
    pic=doc.add_picture(str(img),width=Cm(12));pic._inline.docPr.set('descr','Grupo A: 24 solicitações. Grupo B: 36 solicitações. Barras com rótulos e valores.')
    doc.save(corpus/f'D{i}.docx')
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text=title;s.placeholders[1].text=a+'\n'+b
    s=p.slides.add_slide(p.slide_layouts[5]);s.shapes.title.text='Distribuição das solicitações'
    t=s.shapes.add_table(3,2,Inches(.5),Inches(1.4),Inches(3),Inches(2)).table
    for ri,row in enumerate(rows):
        for ci,value in enumerate(row):t.cell(ri,ci).text=value
    pic=s.shapes.add_picture(str(img),Inches(4),Inches(1.4),width=Inches(5));pic._element.xpath('.//p:cNvPr')[0].set('descr','Grupo A: 24; grupo B: 36 solicitações.')
    p.save(corpus/f'P{i}.pptx')
    pdf=fitz.open();page=pdf.new_page(width=595,height=842)
    page.insert_text((50,60),title,fontsize=18)
    page.insert_textbox((50,90,540,260),a+'\n\n'+b,fontsize=12)
    page.insert_text((50,300),'Grupo      Solicitações\nA                 24\nB                 36',fontsize=12)
    page.insert_image(fitz.Rect(50,380,545,561),filename=str(img));pdf.save(corpus/f'F{i}.pdf');pdf.close()
records=[];spent=0.0
for source in sorted(corpus.iterdir()):
    if source.suffix not in {'.docx','.pptx','.pdf'}:continue
    original=extract(source);cache={}
    for strategy in ['unico','especializadas','base_adaptavel']:
        for repetition in [1,2]:
            target=out/source.stem/strategy/str(repetition);target.mkdir(parents=True,exist_ok=True)
            resultfile=target/'measurement.json'
            if resultfile.exists():
                record=json.loads(resultfile.read_text(encoding='utf-8'));records.append(record);continue
            start=time.perf_counter();usage=[];hits=0;content=copy.deepcopy(original)
            # Duas visualizações nas estratégias de adaptação: original e linguagem simplificada.
            # A primeira é preservada. Os custos da segunda são medidos explicitamente.
            if strategy!='unico':
                for n,block in enumerate(content['blocks']):
                    if block['kind'] not in {'paragraph','heading','image'}:continue
                    ck=fingerprint({'block':block,'model':model_id,'prompt':1})
                    cachefile=out/source.stem/'cache'/f'{ck}.json'
                    if strategy=='base_adaptavel' and cachefile.exists():
                        response=json.loads(cachefile.read_text(encoding='utf-8'));hits+=1
                    else:
                        if not key:raise RuntimeError('Configure OPENROUTER_API_KEY no ambiente local')
                        if spent>.50:raise RuntimeError('Limite conservador de gasto experimental atingido')
                        adapted,u=adapt_block(block,key,model_id,block['kind']!='image',block['kind']=='image',scenarios[int(source.stem[-1])-1][0])
                        response={'block':adapted,'usage':u};usage.append(u);spent+=float(u.get('cost',0) or 0)
                        if strategy=='base_adaptavel':cachefile.parent.mkdir(exist_ok=True);cachefile.write_text(json.dumps(response,ensure_ascii=False),encoding='utf-8')
                    content['blocks'][n]=response['block']
            files,report=export_all(content,target,'Avaliação '+source.stem,{'contrast':True,'pptx':source.suffix=='.pptx'})
            if strategy!='unico':
                original_files,_=export_all(original,target/'original','Versão integral '+source.stem,{'contrast':True})
            (target/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
            comparison=compare(original,content)
            record={'source':source.name,'strategy':strategy,'repetition':repetition,'seconds':time.perf_counter()-start,'calls':len(usage),'cache_hits':hits,'usage':usage,'cost_usd':sum(float(u.get('cost',0) or 0) for u in usage),'bytes':sum(p.stat().st_size for p in target.rglob('*') if p.is_file() and 'lo-profile' not in str(p)),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'model':model_id if strategy!='unico' else None,'temperature':0,'prompt_version':1,'comparison':comparison,'export':report,'files':files,'human_semantic_review':'pendente','source_warnings':original['warnings']}
            resultfile.write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8');records.append(record)
            print(f'{source.name} {strategy} repetição {repetition}: {record["seconds"]:.1f}s; {len(usage)} chamadas',flush=True)
(out/'results.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
summary={'runs':len(records),'cost_usd':sum(r['cost_usd'] for r in records),'model':model_id,'strategies':{s:{'median_seconds':statistics.median(r['seconds'] for r in records if r['strategy']==s),'calls':sum(r['calls'] for r in records if r['strategy']==s),'cache_hits':sum(r['cache_hits'] for r in records if r['strategy']==s)} for s in ['unico','especializadas','base_adaptavel']},'caveat':'Corpus sintético pequeno. Custos incluem tarefas diferentes: uma saída única ou duas visualizações. Cache frio na primeira repetição da base adaptável e quente na segunda. Não mede compreensão, preferência ou equivalência semântica.'}
(out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(summary,ensure_ascii=False))
