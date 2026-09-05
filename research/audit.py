import sys,json,re,statistics,copy
from pathlib import Path
from collections import Counter
from docx import Document
from pypdf import PdfReader
root=Path(__file__).resolve().parents[1];sys.path.insert(0,str(root/'web'))
from studio.content import extract,compare,fingerprint
out=root/'research/evaluation-output'
records=json.loads((out/'results.json').read_text(encoding='utf-8'))
issues=[];pdfs=[];texts=[]
for r in records:
    folder=out/Path(r['source']).stem/r['strategy']/str(r['repetition'])
    d=Document(folder/'documento.docx');text='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
    scenario=int(Path(r['source']).stem[-1]);expected=['30','60','24','36'] if scenario==1 else ['20','15','80','10','24','36']
    missing=[n for n in expected if not re.search(r'(?<!\d)'+n+r'(?!\d)',text)]
    if missing:issues.append({'source':r['source'],'strategy':r['strategy'],'repeat':r['repetition'],'missing_numbers':missing})
    texts.append({'source':r['source'],'strategy':r['strategy'],'repeat':r['repetition'],'text':text})
    if (folder/'documento.pdf').exists():
        p=PdfReader(folder/'documento.pdf');catalog=p.trailer['/Root'];pdfs.append({'source':r['source'],'strategy':r['strategy'],'repeat':r['repetition'],'marked':bool(catalog.get('/MarkInfo',{}).get('/Marked')),'structure':bool(catalog.get('/StructTreeRoot')),'language':str(catalog.get('/Lang','')),'pages':len(p.pages)})
source=extract(root/'research/corpus/D1.docx');updated=copy.deepcopy(source)
for b in updated['blocks']:b['text']=b.get('text','').replace('30 dias','45 dias')
update={'fingerprint_invalidated':fingerprint(source)!=fingerprint(updated),'old_output_against_new_source':compare(updated,source),'scope':'Ensaio de detecção da saída desatualizada. A propagação completa de novas versões não foi medida.'}
report={'runs':len(records),'missing_numeric_facts':issues,'docx_readback_missing':sum(len(r['export']['docx_missing_text']) for r in records),'pdfs':pdfs,'update':update,'semantic_validation':'Leitura assistida por IA; não é revisão humana independente','timing':{s:{str(rep):statistics.median(r['seconds'] for r in records if r['strategy']==s and r['repetition']==rep) for rep in [1,2]} for s in ['unico','especializadas','base_adaptavel']}}
(out/'audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');(out/'delivered-texts.json').write_text(json.dumps(texts,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='pdfs'},ensure_ascii=False));print('PDFs marcados:',sum(p['marked'] and p['structure'] for p in pdfs),'/',len(pdfs))
