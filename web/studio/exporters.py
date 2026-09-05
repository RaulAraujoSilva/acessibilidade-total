import base64
import html
import io
import os
import shutil
import subprocess
import textwrap
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def export_all(content,out,title,options):
    out=Path(out);out.mkdir(parents=True,exist_ok=True)
    d=Document();section=d.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2.5)
    section.left_margin=Cm(3);section.right_margin=Cm(2)
    style=d.styles['Normal'];style.font.name='Arial';style.font.size=Pt(12)
    style.font.color.rgb=RGBColor.from_string('000000' if options.get('contrast') else '132C46')
    style.paragraph_format.line_spacing=1.5
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'pt-BR');style.element.get_or_add_rPr().append(lang)
    d.core_properties.title=title;d.core_properties.language='pt-BR'
    d.add_heading(title,0)
    htmlparts=['<!doctype html><html lang="pt-BR"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>'+html.escape(title)+'</title><style>body{font:1.15rem/1.6 Arial,sans-serif;max-width:70ch;margin:2rem auto;padding:1rem;color:#132c46;background:white}img{max-width:100%;height:auto}td,th{border:1px solid;padding:.5rem}table{border-collapse:collapse}a:focus{outline:3px solid #8a4500}</style><main><h1>'+html.escape(title)+'</h1>']
    expected=[]
    for b in content['blocks']:
        text=b.get('text','');kind=b['kind'];ident=html.escape(b['id'])
        if kind=='heading':
            level=min(3,max(1,b.get('level',1)))
            d.add_heading(text,level);htmlparts.append(f'<h{level+1} id="{ident}">{html.escape(text)}</h{level+1}>');expected.append(text)
        elif kind=='paragraph':
            d.add_paragraph(text);htmlparts.append(f'<p id="{ident}">{html.escape(text)}</p>');expected.append(text)
        elif kind=='table':
            rows=b['rows'];width=max(map(len,rows),default=0)
            if not width:continue
            table=d.add_table(rows=0,cols=width);table.style='Table Grid'
            for i,row in enumerate(rows):
                cells=table.add_row().cells
                for j,val in enumerate(row):cells[j].text=val;expected.append(val)
                if i==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            htmlparts.append(f'<table id="{ident}"><caption>Tabela do {html.escape(b["origin"])}</caption>')
            for i,row in enumerate(rows):
                tag='th scope="col"' if i==0 else 'td';end='th' if i==0 else 'td'
                htmlparts.append('<tr>'+''.join(f'<{tag}>{html.escape(v)}</{end}>' for v in row)+'</tr>')
            htmlparts.append('</table>')
        elif kind=='image':
            alt=b.get('alt','') or 'Descrição pendente de revisão.'
            shape=d.add_picture(io.BytesIO(base64.b64decode(b['image'])),width=Cm(14))
            shape._inline.docPr.set('descr',alt)
            d.add_paragraph(alt,style='Caption')
            htmlparts.append(f'<figure id="{ident}"><img src="data:image/png;base64,{b["image"]}" alt="{html.escape(alt,quote=True)}"><figcaption>{html.escape(alt)}</figcaption></figure>')
    d.add_heading('Relatório de acessibilidade',1)
    d.add_paragraph('Este arquivo foi reconstruído automaticamente e requer revisão. A presença de estrutura e alternativas não comprova compreensão ou conformidade integral.')
    for warning in content.get('warnings',[]):d.add_paragraph(warning)
    docx=out/'documento.docx';d.save(docx)
    htmlparts.append('</main></html>');(out/'documento.html').write_text('\n'.join(htmlparts),encoding='utf-8')
    # Conferir texto realmente serializado no DOCX, além do modelo intermediário.
    read=Document(docx)
    real='\n'.join([p.text for p in read.paragraphs]+[c.text for t in read.tables for row in t.rows for c in row.cells])
    missing=[text[:120] for text in expected if text and text not in real]
    report={'docx_missing_text':missing,'pdf':'indisponível: LibreOffice não encontrado','human_review':'pendente','pdfua':'não verificado','color':'Texto reconstruído não depende de cor; gráficos e imagens originais exigem conferência de rótulos, padrões e descrições.'}
    artifacts=['documento.docx','documento.html']
    if options.get('pptx'):
        _pptx(content,out/'apresentacao.pptx',title,options);artifacts.append('apresentacao.pptx')
        import sys
        legacy=Path(os.environ.get('SKILL_AUDITOR',str(Path(__file__).resolve().parents[2]/'scripts'/'audit_pptx.py')))
        if legacy.exists():
            check=subprocess.run([sys.executable,str(legacy),str(out/'apresentacao.pptx')],capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=60,env={**os.environ,'PYTHONIOENCODING':'utf-8'})
            (out/'auditoria-skill.txt').write_text(check.stdout+check.stderr,encoding='utf-8')
            artifacts.append('auditoria-skill.txt');report['legacy_audit_exit_code']=check.returncode
    soffice=shutil.which('soffice')
    wsl=os.environ.get('LIBREOFFICE_WSL_DISTRO','') if os.name=='nt' else ''
    if soffice or wsl:
        # Perfil isolado evita concorrência entre processos do LibreOffice.
        profile=(out/'lo-profile').resolve().as_uri()
        command=[soffice,f'-env:UserInstallation={profile}','--headless','--convert-to','pdf:writer_pdf_Export:{"PDFUACompliance":{"type":"boolean","value":"true"}}','--outdir',str(out),str(docx)]
        if wsl:
            def lp(p):
                value=Path(p).resolve().as_posix().removeprefix('//?/');return '/mnt/'+value[0].lower()+value[2:]
            import uuid
            command=['wsl','-d',wsl,'--exec','libreoffice','-env:UserInstallation=file:///tmp/accessible-lo-'+uuid.uuid4().hex,'--headless','--convert-to','pdf:writer_pdf_Export:{"PDFUACompliance":{"type":"boolean","value":"true"}}','--outdir',lp(out),lp(docx)]
        try:result=subprocess.run(command,capture_output=True,timeout=120)
        except subprocess.TimeoutExpired:
            report['pdf']='tempo limite na exportação; DOCX preservado'
            return artifacts,report
        pdf=out/'documento.pdf'
        if result.returncode==0 and pdf.exists():
            import fitz
            with fitz.open(pdf) as p:report['pdf']={'pages':len(p),'searchable':any(x.get_text().strip() for x in p)}
            artifacts.append('documento.pdf')
            vera=shutil.which('verapdf')
            if vera:
                v=subprocess.run([vera,'--format','xml',str(pdf)],capture_output=True,text=True,timeout=120)
                (out/'verapdf.xml').write_text(v.stdout,encoding='utf-8');artifacts.append('verapdf.xml')
                report['pdfua']='validação automática sem reprovação' if v.returncode==0 and 'isCompliant="true"' in v.stdout and 'isCompliant="false"' not in v.stdout else 'reprovado ou indeterminado'
        else:report['pdf']='falha na exportação; DOCX preservado'
    return artifacts,report

def _pptx(content,path,title,options):
    from pptx import Presentation
    from pptx.util import Inches,Pt as PPt
    from pptx.dml.color import RGBColor as PRGB
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    p.core_properties.title=title;p.core_properties.language='pt-BR'
    heading=title
    for b in content['blocks']:
        if b['kind']=='heading':heading=b['text'];continue
        if b['kind']=='image':
            slide=p.slides.add_slide(p.slide_layouts[5]);slide.shapes.title.text=heading[:140]
            from PIL import Image
            data=base64.b64decode(b['image'])
            with Image.open(io.BytesIO(data)) as im:ratio=im.width/im.height
            size={'width':Inches(11)} if ratio>11/4.8 else {'height':Inches(4.8)}
            pic=slide.shapes.add_picture(io.BytesIO(data),Inches(1),Inches(1.6),**size)
            pic._element.xpath('.//p:cNvPr')[0].set('descr',b.get('alt','') or 'Descrição pendente de revisão.')
            slide.notes_slide.notes_text_frame.text=b.get('alt','');continue
        if b['kind']=='table' and b.get('rows'):
            rows=b['rows'];header=rows[0];cols=max(map(len,rows))
            chunks=[rows[i:i+6] for i in range(1,len(rows),6)] or [[]]
            for chunk in chunks:
                slide=p.slides.add_slide(p.slide_layouts[5]);slide.shapes.title.text=heading[:140]
                table_shape=slide.shapes.add_table(len(chunk)+1,cols,Inches(.7),Inches(1.7),Inches(11.9),Inches(min(4.8,(len(chunk)+1)*.65)))
                table_shape._element.xpath('.//p:cNvPr')[0].set('descr','Tabela: '+', '.join(header)+'. '+ '; '.join(', '.join(row) for row in chunk))
                table=table_shape.table
                table.first_row=True
                for ri,row in enumerate([header]+chunk):
                    for ci,value in enumerate(row):
                        table.cell(ri,ci).text=value
                        cell=table.cell(ri,ci);cell.fill.solid();cell.fill.fore_color.rgb=PRGB.from_string('132C46' if ri==0 else 'FFFFFF')
                        for para in cell.text_frame.paragraphs:
                            para.font.size=PPt(18);para.font.color.rgb=PRGB.from_string('FFFFFF' if ri==0 else '132C46')
                            for run in para.runs:
                                run._r.get_or_add_rPr().set('lang','pt-BR');run.font.color.rgb=PRGB.from_string('FFFFFF' if ri==0 else '132C46');run.font.size=PPt(18)
                slide.notes_slide.notes_text_frame.text=f'Origem: {b["origin"]}; bloco {b["id"]}. Conferir legibilidade e ordem da tabela.'
            continue
        text=b.get('text','')
        if b['kind']=='table':text='\n'.join(' | '.join(row) for row in b['rows'])
        for n,part in enumerate(textwrap.wrap(text,width=420,replace_whitespace=False,drop_whitespace=False) or ['']):
            slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text=heading[:130]+(f' ({n+1})' if n else '')
            body=slide.placeholders[1];body.text=part
            for para in body.text_frame.paragraphs:
                para.font.size=PPt(24);para.font.color.rgb=PRGB.from_string('000000' if options.get('contrast') else '132C46')
            slide.notes_slide.notes_text_frame.text=f'Origem: {b["origin"]}; bloco {b["id"]}. '+('Tabela também disponível em estrutura semântica no DOCX/HTML.' if b['kind']=='table' else '')
    if not len(p.slides):
        slide=p.slides.add_slide(p.slide_layouts[5]);slide.shapes.title.text=heading
    for index,slide in enumerate(p.slides,1):
        title_shape=slide.shapes.title
        title_shape.text=title_shape.text+f' — {index} de {len(p.slides)}'
        title_shape.left=Inches(.6);title_shape.top=Inches(.6);title_shape.width=Inches(12.1);title_shape.height=Inches(.75)
        for para in title_shape.text_frame.paragraphs:para.font.size=PPt(28)
        for number,shape in enumerate(slide.shapes,1):
            shape.name=f'Elemento {number} do slide {index}'
            if shape.has_text_frame and shape!=title_shape:
                shape.left=Inches(.7);shape.top=Inches(1.7);shape.width=Inches(11.9);shape.height=Inches(5)
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:run._r.get_or_add_rPr().set('lang','pt-BR')
    p.save(path)
