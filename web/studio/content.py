"""Modelo documental independente da interface web e do provedor de IA."""
import base64
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from PIL import Image

SUPPORTED={'.pptx','.docx','.pdf','.txt','.md'}

def fingerprint(content):
    return hashlib.sha256(json.dumps(content,ensure_ascii=False,sort_keys=True).encode()).hexdigest()

def _image(blob):
    with Image.open(io.BytesIO(blob)) as im:
        if im.width*im.height>25_000_000: raise ValueError('Imagem muito grande para o protótipo.')
        im.thumbnail((2000,2000));out=io.BytesIO();im.convert('RGB').save(out,format='PNG')
    return base64.b64encode(out.getvalue()).decode()

def _zip_check(path):
    with zipfile.ZipFile(path) as z:
        if len(z.infolist())>10000 or sum(i.file_size for i in z.infolist())>250*1024*1024:
            raise ValueError('Documento compactado excede o limite de processamento.')
        if any('vbaProject' in i.filename for i in z.infolist()):
            raise ValueError('Documentos com macros não são aceitos.')

def extract(path):
    path=Path(path);ext=path.suffix.lower()
    if ext not in SUPPORTED: raise ValueError('Use PPTX, DOCX, PDF digital, TXT ou Markdown.')
    if path.stat().st_size>50*1024*1024:raise ValueError('Limite de 50 MB por documento.')
    blocks=[];warnings=[]
    def add(kind,text='',location='',**extra):
        blocks.append({'id':f'b{len(blocks)+1:04d}','kind':kind,'text':text,'origin':location,**extra})
    if ext in {'.docx','.pptx'}:_zip_check(path)
    if ext=='.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        d=Document(path)
        for n,el in enumerate(d.element.body):
            origin=f'elemento {n+1}'
            if el.tag==qn('w:p'):
                p=Paragraph(el,d)
                if p.text.strip():
                    name=p.style.name if p.style else ''
                    heading=re.search(r'(?:Heading|Título)\s*(\d)',name)
                    add('heading' if heading else 'paragraph',p.text,origin,level=int(heading.group(1)) if heading else 0)
                for blip in el.xpath('.//a:blip'):
                    rid=blip.get(qn('r:embed'))
                    if rid and rid in d.part.related_parts:
                        props=el.xpath('.//wp:docPr')
                        alt=props[0].get('descr','') if props else ''
                        try:add('image',alt,origin,image=_image(d.part.related_parts[rid].blob),alt=alt)
                        except (ValueError,OSError):warnings.append(f'{origin}: imagem não recuperada.')
                if el.xpath('.//m:oMath'):warnings.append(f'{origin}: equação exige transcrição e revisão.')
            elif el.tag==qn('w:tbl'):
                t=Table(el,d);rows=[[c.text for c in row.cells] for row in t.rows]
                add('table','',origin,rows=rows)
                if el.xpath('.//w:gridSpan') or el.xpath('.//w:vMerge'):warnings.append(f'{origin}: tabela com mesclagem; conferir reconstrução.')
        if d.part.footnotes_part if hasattr(d.part,'footnotes_part') else False:
            warnings.append('Notas de rodapé requerem conferência.')
        warnings.append('Conferir cabeçalhos, rodapés, notas e objetos flutuantes; não são reconstruídos automaticamente.')
    elif ext=='.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        p=Presentation(path)
        if len(p.slides)>100:raise ValueError('Limite de 100 slides.')
        for i,slide in enumerate(p.slides):
            origin=f'slide {i+1}'
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    is_title=shape==slide.shapes.title
                    add('heading' if is_title else 'paragraph',shape.text,origin,level=1 if is_title else 0)
                elif shape.has_table:
                    add('table','',origin,rows=[[c.text for c in row.cells] for row in shape.table.rows])
                elif shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
                    props=shape._element.xpath('.//p:cNvPr');alt=props[0].get('descr','') if props else ''
                    try:add('image',alt,origin,image=_image(shape.image.blob),alt=alt)
                    except (ValueError,OSError):warnings.append(f'{origin}: imagem não recuperada.')
                elif shape.shape_type==MSO_SHAPE_TYPE.GROUP or getattr(shape,'has_chart',False):
                    warnings.append(f'{origin}: grupo/gráfico exige revisão e descrição.')
            if slide.has_notes_slide:
                notes=slide.notes_slide.notes_text_frame.text.strip()
                if notes:add('paragraph',notes,origin+' notas',role='notes')
    elif ext=='.pdf':
        import fitz
        with fitz.open(path) as doc:
            if doc.is_encrypted:raise ValueError('PDF protegido; envie uma cópia autorizada sem senha.')
            if len(doc)>100:raise ValueError('Limite de 100 páginas.')
            for i,page in enumerate(doc):
                text=page.get_text(sort=True).strip()
                if not text:
                    warnings.append(f'Página {i+1} sem texto: OCR não disponível.')
                    continue
                add('heading',f'Página {i+1}',f'página {i+1}',level=1,generated_label=True)
                for part in re.split(r'\n\s*\n',text):add('paragraph',part,f'página {i+1}')
                for im in page.get_images(full=True):
                    try:
                        blob=doc.extract_image(im[0])['image'];add('image','',f'página {i+1}',image=_image(blob),alt='')
                    except (ValueError,OSError):warnings.append(f'Página {i+1}: imagem não recuperada.')
            warnings.append('PDF: tabelas, colunas, fórmulas e ordem semântica exigem revisão; extração geométrica não comprova estrutura.')
    else:
        text=path.read_text(encoding='utf-8-sig')
        for n,part in enumerate(re.split(r'\n\s*\n',text)):
            if not part.strip():continue
            match=re.match(r'^(#{1,3})\s+(.+)',part)
            add('heading' if match else 'paragraph',match.group(2) if match else part,f'bloco {n+1}',level=len(match.group(1)) if match else 0)
            if match and part[match.end():].strip():add('paragraph',part[match.end():].strip(),f'bloco {n+1}, após título')
    if not blocks or not any(b['text'] for b in blocks):raise ValueError('Nenhum texto recuperado. Digitalizações não são suportadas nesta versão.')
    if len(blocks)>1000:raise ValueError('Limite de 1000 blocos por documento.')
    return {'schema_version':1,'source_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'blocks':blocks,'warnings':warnings}

def compare(source,adapted):
    """Verifica a saída real. Igualdade estrutural não é equivalência semântica."""
    base={b['id']:b for b in source['blocks']};other={b['id']:b for b in adapted['blocks']}
    report={'missing_blocks':sorted(base.keys()-other.keys()),'added_blocks':sorted(other.keys()-base.keys()),'changed':[],'numeric_losses':[],'semantic_equivalence':'não verificada; requer revisão humana'}
    for ident,b in base.items():
        if ident not in other:continue
        a=other[ident]
        if b.get('text')!=a.get('text'):report['changed'].append(ident)
        nums=lambda t: set(re.findall(r'\b\d+(?:[.,]\d+)*%?',t))
        lost=nums(b.get('text',''))-nums(a.get('text',''))
        if lost:report['numeric_losses'].append({'id':ident,'values':sorted(lost)})
        if b.get('rows')!=a.get('rows') or b.get('image')!=a.get('image'):
            report.setdefault('changed_structures',[]).append(ident)
    report['requires_review']=True
    return report
