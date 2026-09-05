import copy
import io
import json
import tempfile
import uuid
from pathlib import Path
from unittest.mock import patch,Mock
from cryptography.fernet import Fernet
from django.test import TransactionTestCase,override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core import mail
from django.urls import reverse
from django.utils import timezone
from datetime import timedelta
from docx import Document as Word
from .models import User,Document,ProviderKey,StageResult
from .security import seal,unseal
from .content import extract,compare
from .providers import adapt_block,ProviderError,TemporaryProviderError
from .tasks import execute
from .exporters import export_all

@override_settings(DEBUG=True,DISPATCH_MODE='database',EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',FERNET_KEY=Fernet.generate_key().decode(),SECURE_SSL_REDIRECT=False)
class WorkflowTests(TransactionTestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.settings_override=override_settings(MEDIA_ROOT=Path(self.tmp.name));self.settings_override.enable();self.addCleanup(self.settings_override.disable)
        self.user=User.objects.create_user(username='one@example.test',email='one@example.test',password='Test-long-password-1957')
        self.other=User.objects.create_user(username='two@example.test',email='two@example.test',password='Test-long-password-2789')
        self.client.force_login(self.user)
    def document(self):
        source=SimpleUploadedFile('fonte.md','# Condições\n\nO prazo é de 30 dias, exceto em recurso.'.encode())
        doc=Document.objects.create(owner=self.user,title='Condições',source=source,original_name='fonte.md')
        doc.content=extract(doc.source.path);doc.options={'contrast':True};doc.run_id=uuid.uuid4();doc.status='queued';doc.save();return doc
    def test_upload_extract_does_not_call_ai(self):
        with patch('studio.providers.requests.post') as post:
            r=self.client.post(reverse('upload'),{'title':'Teste','source':SimpleUploadedFile('t.txt','A informação deve ser preservada.'.encode())})
        self.assertEqual(r.status_code,302);post.assert_not_called();self.assertEqual(Document.objects.get().status,'created')
    def test_markdown_heading_does_not_discard_following_line(self):
        path=Path(self.tmp.name)/'t.md';path.write_text('# Condição\nNão generalizar o resultado de 10 participantes.',encoding='utf-8')
        content=extract(path)
        self.assertEqual(content['blocks'][0]['kind'],'heading')
        self.assertIn('Não generalizar',content['blocks'][1]['text'])
    def test_other_account_cannot_read_or_download(self):
        doc=self.document();self.client.force_login(self.other)
        for url in [reverse('detail',args=[doc.pk]),reverse('status',args=[doc.pk]),reverse('download',args=[doc.pk,0])]:self.assertEqual(self.client.get(url).status_code,404)
        self.assertEqual(self.client.post(reverse('cancel',args=[doc.pk])).status_code,404)
    def test_key_encrypted_at_rest_and_not_echoed(self):
        value='sk-or-v1-not-a-real-secret-123456'
        self.client.post(reverse('keys'),{'key':value});record=ProviderKey.objects.get()
        self.assertNotIn(value,record.ciphertext);self.assertEqual(unseal(record.ciphertext),value)
        self.assertNotContains(self.client.get(reverse('keys')),value)
    def test_pipeline_generates_real_files_and_requires_review(self):
        doc=self.document();execute(str(doc.pk),str(doc.run_id));doc.refresh_from_db()
        self.assertEqual(doc.status,'review');self.assertEqual(doc.report['docx_missing_text'],[])
        self.assertTrue(doc.report['comparison']['requires_review'])
        response=self.client.get(reverse('download',args=[doc.pk,0]))
        self.assertEqual(response.status_code,200);response.close()
    def test_cancelled_task_cannot_resurrect(self):
        doc=self.document();run=doc.run_id;self.client.post(reverse('cancel',args=[doc.pk]));execute(str(doc.pk),str(run));doc.refresh_from_db()
        self.assertEqual(doc.status,'cancelled');self.assertEqual(doc.artifacts,[])
    def test_stale_lease_can_resume(self):
        doc=self.document();doc.status='running';doc.heartbeat=timezone.now()-timedelta(minutes=10);doc.save()
        execute(str(doc.pk),str(doc.run_id));doc.refresh_from_db();self.assertEqual(doc.status,'review')
    def test_completed_blocks_reused_after_retry(self):
        doc=self.document();doc.options={'simplify':True};doc.model_id='test/model';doc.save()
        ProviderKey.objects.create(owner=self.user,ciphertext=seal('sk-or-fake-key'))
        def fail_second(block,*a,**kw):
            if block['kind']=='paragraph':raise TemporaryProviderError('temporário')
            return block,{'total_tokens':4}
        with patch('studio.tasks.adapt_block',side_effect=fail_second):
            with self.assertRaises(TemporaryProviderError):execute(str(doc.pk),str(doc.run_id))
        self.assertEqual(StageResult.objects.count(),1)
        with patch('studio.tasks.adapt_block',side_effect=lambda b,*a,**k:(b,{})) as api:execute(str(doc.pk),str(doc.run_id))
        self.assertEqual(api.call_count,1)
    def test_actual_text_loss_detected_even_with_identical_ids(self):
        source={'blocks':[{'id':'x','kind':'paragraph','text':'O prazo é de 30 dias, exceto em recurso.'}]}
        changed=copy.deepcopy(source);changed['blocks'][0]['text']='Existe um prazo.'
        report=compare(source,changed)
        self.assertEqual(report['changed'],['x']);self.assertTrue(report['numeric_losses'])
    def test_empty_pdf_or_bad_format_is_not_accepted(self):
        p=Path(self.tmp.name)/'x.exe';p.write_bytes(b'bad')
        with self.assertRaises(ValueError):extract(p)
    def test_docx_table_and_image_are_preserved(self):
        from PIL import Image
        p=Path(self.tmp.name)/'x.docx';d=Word();d.add_heading('Tabela',1);t=d.add_table(rows=2,cols=2)
        t.cell(0,0).text='Item';t.cell(0,1).text='Valor';t.cell(1,0).text='Prazo';t.cell(1,1).text='30'
        image=io.BytesIO();Image.new('RGB',(20,20),'blue').save(image,'PNG');image.seek(0);d.add_picture(image);d.save(p)
        result=extract(p);self.assertIn('table',[b['kind'] for b in result['blocks']]);self.assertIn('image',[b['kind'] for b in result['blocks']])
        artifacts,report=export_all(result,Path(self.tmp.name)/'out','Teste',{})
        self.assertFalse(report['docx_missing_text'])
    def test_signup_requires_email_confirmation(self):
        self.client.logout();r=self.client.post(reverse('signup'),{'email':'new@example.test','password1':'Strong-secret-1987','password2':'Strong-secret-1987'})
        self.assertEqual(r.status_code,200);self.assertFalse(User.objects.get(email='new@example.test').is_active);self.assertEqual(len(mail.outbox),1)
    def test_provider_errors_are_safe_and_classified(self):
        for status,error in [(401,ProviderError),(402,ProviderError),(429,TemporaryProviderError)]:
            with patch('studio.providers.requests.post',return_value=Mock(status_code=status,ok=False)):
                with self.assertRaises(error) as cm:adapt_block({'kind':'paragraph','text':'Teste'},'private-token','m',True)
                self.assertNotIn('private-token',str(cm.exception))
    def test_public_mutations_require_csrf(self):
        from django.test import Client
        secure=Client(enforce_csrf_checks=True);secure.force_login(self.user)
        self.assertEqual(secure.post(reverse('delete_key')).status_code,403)
    def test_download_path_traversal_is_rejected(self):
        doc=self.document();doc.status='review';doc.artifacts=['../../../../secret.txt'];doc.save()
        self.assertEqual(self.client.get(reverse('download',args=[doc.pk,0])).status_code,404)
    def test_approve_requires_explicit_review(self):
        doc=self.document();doc.status='review';doc.save()
        self.assertEqual(self.client.post(reverse('approve',args=[doc.pk])).status_code,400)
    def test_libras_missing_engine_remains_explicit(self):
        from .libras import render,LibrasUnavailable
        with patch.dict('os.environ',{'VLIBRAS_RENDERER':'','VLIBRAS_BUNDLES':''}):
            with self.assertRaises(LibrasUnavailable):render('Olá',Path(self.tmp.name))

    def test_libras_live_lease_prevents_duplicate_renderer(self):
        from .tasks import render_libras
        doc=self.document();doc.status='review';doc.report={'libras':{'status':'aguardando worker'}}
        doc.stage='Renderizando Libras';doc.heartbeat=timezone.now();doc.save()
        with patch('studio.tasks.render') as renderer:
            render_libras.apply(args=[str(doc.pk),str(doc.run_id)]).get()
        renderer.assert_not_called()

    def test_libras_storage_failure_is_visible_not_stuck(self):
        doc=self.document();doc.options={'libras':True};doc.save()
        with patch('studio.tasks.render',side_effect=OSError('private-path')):
            execute(str(doc.pk),str(doc.run_id))
        doc.refresh_from_db()
        self.assertEqual(doc.report['libras']['status'],'indisponível')
        self.assertNotIn('private-path',doc.report['libras']['error'])

    def test_libras_download_rejects_html_and_unsafe_filename(self):
        from .libras import prepare_bundles
        from unittest.mock import MagicMock
        target=Path(self.tmp.name)/'bundles';(target/'BR').mkdir(parents=True)
        for token in 'ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789':(target/'BR'/token).write_bytes(b'installed')
        response=MagicMock();response.__enter__.return_value=response;response.ok=True
        response.iter_content.return_value=[b'<html>Error page</html>']
        with patch.dict('os.environ',{'VLIBRAS_FETCH_BUNDLES':'1'}),patch('studio.libras.requests.get',return_value=response) as fetch:
            missing=prepare_bundles('SINAL ../escape',target)
        self.assertIn('SINAL',missing);self.assertIn('../escape',missing)
        self.assertFalse((target/'BR'/'SINAL').exists());self.assertEqual(fetch.call_count,1)

    def test_recovery_republishes_stale_libras_only(self):
        from django.core.management import call_command
        doc=self.document();doc.status='review';doc.report={'libras':{'status':'aguardando worker'}}
        doc.stage='Renderizando Libras';doc.heartbeat=timezone.now();doc.save()
        with patch('studio.management.commands.recover_jobs.render_libras.delay') as dispatch:
            call_command('recover_jobs',stdout=io.StringIO());dispatch.assert_not_called()
            doc.heartbeat=timezone.now()-timedelta(minutes=20);doc.save()
            call_command('recover_jobs',stdout=io.StringIO());dispatch.assert_called_once_with(str(doc.pk),str(doc.run_id))

    def test_gmail_backend_sends_mime_without_exposing_credentials(self):
        import base64
        from django.core.mail import EmailMessage
        from .mail import GmailBackend
        token=Mock();token.json.return_value={'access_token':'temporary-secret'}
        sent=Mock()
        with patch('studio.mail.requests.post',side_effect=[token,sent]) as post:
            count=GmailBackend().send_messages([EmailMessage('Confirmação','Texto de teste','sender@example.test',['recipient@example.test'])])
        self.assertEqual(count,1)
        raw=post.call_args_list[1].kwargs['json']['raw']
        mime=base64.urlsafe_b64decode(raw).decode()
        self.assertIn('recipient@example.test',mime)
        self.assertNotIn('temporary-secret',mime)

    def test_gmail_backend_error_is_sanitized(self):
        import requests
        from django.core.mail import EmailMessage
        from .mail import GmailBackend
        with patch('studio.mail.requests.post',side_effect=requests.RequestException('sensitive-provider-response')):
            with self.assertRaises(OSError) as error:
                GmailBackend().send_messages([EmailMessage('Teste','Teste')])
        self.assertNotIn('sensitive-provider-response',str(error.exception))

    def test_retry_libras_preserves_documents_and_failure_history(self):
        from django.core.management import call_command,CommandError
        doc=self.document();doc.status='review';doc.artifacts=['documento.docx']
        doc.report={'libras':{'status':'indisponível','error':'Serviço indisponível'}};doc.save()
        with patch('studio.management.commands.retry_libras.render_libras.delay') as dispatch:
            call_command('retry_libras',str(doc.pk),stdout=io.StringIO())
            dispatch.assert_called_once_with(str(doc.pk),str(doc.run_id))
        doc.refresh_from_db()
        self.assertEqual(doc.artifacts,['documento.docx'])
        self.assertEqual(doc.report['libras_attempts'][0]['status'],'indisponível')
        with self.assertRaises(CommandError):call_command('retry_libras',str(doc.pk),stdout=io.StringIO())

    def test_proxy_client_ip_requires_explicit_trust(self):
        from django.test import RequestFactory
        from .security import limited
        request=RequestFactory().get('/',REMOTE_ADDR='10.0.0.2',HTTP_X_REAL_IP='198.51.100.1')
        with override_settings(TRUST_PROXY_CLIENT_IP=False):
            self.assertFalse(limited(request,'untrusted',1))
            request.META['HTTP_X_REAL_IP']='198.51.100.2'
            self.assertTrue(limited(request,'untrusted',1))
        with override_settings(TRUST_PROXY_CLIENT_IP=True):
            self.assertFalse(limited(request,'trusted',1))
            request.META['HTTP_X_REAL_IP']='198.51.100.3'
            self.assertFalse(limited(request,'trusted',1))
